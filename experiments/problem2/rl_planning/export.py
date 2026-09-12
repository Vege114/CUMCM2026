"""Export archived exp007 RL decisions through the original question workbook.

The worksheet author is the repository's existing artifact-tool template author.
Numerical checks below are independent of PPO, its runner, and execution code.
"""

import argparse
import hashlib
import json
import subprocess
import time
from datetime import UTC, datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "data/results/exp007"
REPORT = ROOT / "reports/experiments/exp007"
WORK = ROOT / ".work/exp007"
CASE = "regularized_42"
SPECIFIED = ("2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21")
ETA = float(np.sqrt(0.9))
SOC_MIN, SOC_MAX, AC_LIMIT = 1200.0, 10800.0, 5000.0 / 6


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write_json(path, obj):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2, allow_nan=False) + "\n")


def read_npz(path):
    with np.load(path) as arrays:
        return {name: arrays[name].copy() for name in arrays.files}


def clock(slot):
    return f"{slot // 6:02d}:{slot % 6 * 10:02d}"


def periods(energy):
    active = np.asarray(energy) > 1e-6
    edges = np.diff(np.r_[False, active, False].astype(int))
    return [(int(a), int(b), float(np.sum(energy[a:b])))
            for a, b in zip(np.flatnonzero(edges == 1), np.flatnonzero(edges == -1))]


def independent_check(d, days=334, check_projection=True, check_greedy=False):
    """Check accounting identities from arrays, without replaying the controller."""
    for name in ("original", "final", "charge", "discharge", "emergency", "surplus", "price"):
        assert d[name].shape == (days, 144), (name, d[name].shape)
        assert np.isfinite(d[name]).all() and d[name].min() >= -1e-6, name
    assert d["actual"].shape[:2] == (days, 144) and d["actual"].shape[-1] >= 2
    assert np.isfinite(d["actual"]).all() and d["actual"].min() >= 0
    assert d["states"].shape == (days, 145) and np.isfinite(d["states"]).all()
    assert d["fees"].shape == (days, 144, 4) and np.isfinite(d["fees"]).all()
    np.testing.assert_allclose(d["original"], d["final"], rtol=0, atol=1e-8)
    np.testing.assert_allclose(d["states"][:-1, -1], d["states"][1:, 0], rtol=0, atol=1e-6)
    assert SOC_MIN - 1e-6 <= d["states"].min() <= d["states"].max() <= SOC_MAX + 1e-6
    assert max(d["charge"].max(), d["discharge"].max()) <= AC_LIMIT + 1e-6
    simultaneous = int(np.count_nonzero((d["charge"] > 1e-6) & (d["discharge"] > 1e-6)))
    emergency_charge = int(np.count_nonzero((d["charge"] > 1e-6) & (d["emergency"] > 1e-6)))
    assert simultaneous == emergency_charge == 0
    transition = np.diff(d["states"], axis=1) - ETA * d["charge"] + d["discharge"] / ETA
    net = d["final"] + (d["actual"][:, :, 1] - d["actual"][:, :, 0]) / 6
    balance = net + d["discharge"] + d["emergency"] - d["charge"] - d["surplus"]
    assert np.abs(transition).max() < 1e-6 and np.abs(balance).max() < 1e-6
    assert not ((d["emergency"] > 1e-6) & (d["surplus"] > 1e-6)).any()
    # Fixed midnight purchases: Q2 has no adjustment fees or terminal credits.
    fees = np.zeros((days, 144, 4))
    fees[:, :, 0] = d["original"] * d["price"]
    fees[:, :, 3] = 5 * d["emergency"] * d["price"]
    np.testing.assert_allclose(d["fees"], fees, rtol=0, atol=1e-6)
    projected = "not evaluated for warmup"
    if check_projection:
        # The intended requests may be clipped, and are never forced to greedy.
        assert "intended_charge" in d and "intended_discharge" in d
        for name in ("intended_charge", "intended_discharge"):
            assert d[name].shape == (days, 144)
            assert np.isfinite(d[name]).all() and d[name].min() >= -1e-6
            assert d[name].max() <= AC_LIMIT + 1e-6
        assert not ((d["intended_charge"] > 1e-6) & (d["intended_discharge"] > 1e-6)).any()
        if "intended_states" in d:
            assert d["intended_states"].shape == (days, 145)
            assert np.isfinite(d["intended_states"]).all()
            assert SOC_MIN - 1e-6 <= d["intended_states"].min() <= d["intended_states"].max() <= SOC_MAX + 1e-6
            np.testing.assert_allclose(d["intended_states"][:, 0], d["states"][:, 0], rtol=0, atol=1e-6)
            np.testing.assert_allclose(np.diff(d["intended_states"], axis=1),
                                       ETA * d["intended_charge"] - d["intended_discharge"] / ETA,
                                       rtol=0, atol=1e-6)
        expected_c = np.minimum.reduce((d["intended_charge"], np.maximum(net, 0),
                                       np.full_like(net, AC_LIMIT),
                                       np.maximum((SOC_MAX-d["states"][:, :-1])/ETA, 0)))
        expected_d = np.minimum.reduce((d["intended_discharge"], np.maximum(-net, 0),
                                       np.full_like(net, AC_LIMIT),
                                       np.maximum((d["states"][:, :-1]-SOC_MIN)*ETA, 0)))
        expected_c[expected_c < 2.0] = 0
        expected_d[expected_d < 2.0] = 0
        np.testing.assert_allclose(d["charge"], expected_c, rtol=0, atol=1e-6)
        np.testing.assert_allclose(d["discharge"], expected_d, rtol=0, atol=1e-6)
        projected = "passed; vectorized request, balance, SOC and power clipping with 2 kWh deadband"
    if check_greedy:
        expected_c = np.minimum.reduce((np.maximum(net, 0), np.full_like(net, AC_LIMIT),
                                       np.maximum((SOC_MAX-d["states"][:, :-1])/ETA, 0)))
        expected_d = np.minimum.reduce((np.maximum(-net, 0), np.full_like(net, AC_LIMIT),
                                       np.maximum((d["states"][:, :-1]-SOC_MIN)*ETA, 0)))
        np.testing.assert_allclose(d["charge"], expected_c, rtol=0, atol=1e-6)
        np.testing.assert_allclose(d["discharge"], expected_d, rtol=0, atol=1e-6)
        projected = "not applicable; independently verified causal greedy execution with no deadband"
    nonidle = np.sign(d["charge"].ravel() - d["discharge"].ravel())
    nonidle = nonidle[nonidle != 0]
    return {"days": days, "intervals": days * 144, "violations": 0,
            "max_balance_error_kwh": float(np.abs(balance).max()),
            "max_state_error_kwh": float(np.abs(transition).max()),
            "minimum_soc_kwh": float(d["states"].min()), "maximum_soc_kwh": float(d["states"].max()),
            "max_charge_kwh_per_slot": float(d["charge"].max()),
            "max_discharge_kwh_per_slot": float(d["discharge"].max()),
            "simultaneous_slots": simultaneous, "emergency_charging_slots": emergency_charge,
            "fixed_midnight_purchase_plan": "passed", "projection_check": projected,
            "total_cost_yuan": float(fees.sum()), "planned_cost_yuan": float(fees[:, :, 0].sum()),
            "emergency_cost_yuan": float(fees[:, :, 3].sum()),
            "charge_kwh": float(d["charge"].sum()), "discharge_kwh": float(d["discharge"].sum()),
            "equivalent_full_cycles": float((ETA * d["charge"].sum() + d["discharge"].sum()/ETA)/24000),
            "direction_reversals_within_evaluation": int(np.count_nonzero(nonidle[1:] * nonidle[:-1] == -1)),
            "initial_soc_kwh": float(d["states"][0, 0]), "final_soc_kwh": float(d["states"][-1, -1])}


def prepare():
    WORK.mkdir(parents=True, exist_ok=True)
    table_dir = REPORT / "specified_tables"
    table_dir.mkdir(parents=True, exist_ok=True)
    d = read_npz(OUT / CASE / "dispatch_2.npz")
    check = independent_check(d, check_projection=True)
    daily = pd.read_csv(OUT / CASE / "daily.csv")
    assert len(daily) == 334
    if pd.api.types.is_numeric_dtype(daily["day"]):
        dates = (pd.Timestamp("2025-01-01") + pd.to_timedelta(daily["day"], unit="D")).dt.strftime("%Y-%m-%d").tolist()
    else:
        dates = pd.to_datetime(daily["day"]).dt.strftime("%Y-%m-%d").tolist()
    expected_dates = pd.date_range("2025-02-01", "2025-12-31").strftime("%Y-%m-%d").tolist()
    assert dates == expected_dates
    np.testing.assert_allclose(daily["total_cost"], d["fees"].sum((1, 2)), rtol=0, atol=1e-6)
    # Re-read source tables independently, avoiding the shared Data class.
    source_hashes = {}
    actual_arrays = []
    for name in ("附件2_小区负载.csv", "附件2_光伏发电实际功率.csv"):
        path = ROOT / "data/raw" / name
        source = pd.read_csv(path)
        assert source.shape == (365, 145)
        assert pd.to_datetime(source.iloc[31:, 0]).dt.strftime("%Y-%m-%d").tolist() == dates
        actual_arrays.append(source.iloc[31:, 1:].to_numpy(float))
        source_hashes[name] = sha256(path)
    np.testing.assert_allclose(d["actual"][:, :, :2], np.stack(actual_arrays, -1), rtol=0, atol=1e-8)
    price_path = ROOT / "data/raw/附件1.csv"
    tariff = pd.read_csv(price_path)["电价"].to_numpy(float)
    np.testing.assert_allclose(d["price"], np.broadcast_to(tariff, (334, 144)), rtol=0, atol=1e-10)
    source_hashes[price_path.name] = sha256(price_path)
    warm_path = ROOT / "data/results/exp003/warmup_2.npz"
    warm = read_npz(warm_path)
    warm_check = independent_check(warm, days=31, check_projection=False)
    np.testing.assert_allclose(warm["states"][0, 0], 6000.0, rtol=0, atol=1e-8)
    np.testing.assert_allclose(d["states"][0, 0], warm["states"][-1, -1], rtol=0, atol=1e-8)
    battery, emergency, all_battery, all_emergency = [], [], [], []
    summaries, intervals, specified_battery, specified_emergency = [], [], [], []
    role = "预先确定的 regularized_42 主组" if CASE == "regularized_42" else f"{CASE} 费用优先消融组"
    markdown = [f"# exp007 {role}指定四日结果", f"\n电量与储电量为 kWh，费用为元。采用{role}（复用无季节预测种子 42；PPO 强化学习规划）。表 2 为真实负载下的执行量。全天费用包含计划购电费和 5 倍电价的紧急购电费。"]
    for i, date in enumerate(dates):
        today_battery = []
        for block in range(6):
            start, stop = block * 24, (block + 1) * 24
            c, discharge = float(d["charge"][i, start:stop].sum()), float(d["discharge"][i, start:stop].sum())
            period = f"{block * 4}:00-{(block + 1) * 4}:00"
            battery.append([date if block == 0 else None, period, c, discharge,
                            "0:00" if block == 0 else "24:00" if block == 1 else None,
                            float(d["states"][i, 0]) if block == 0 else float(d["states"][i, -1]) if block == 1 else None])
            record = {"date": date, "period": period, "charge_kwh": c, "discharge_kwh": discharge,
                      "initial_soc_kwh": float(d["states"][i, 0]), "final_soc_kwh": float(d["states"][i, -1])}
            all_battery.append(record)
            today_battery.append(record)
        events = periods(d["emergency"][i])
        np.testing.assert_allclose(sum(e[2] for e in events), d["emergency"][i].sum(), rtol=0, atol=1e-5)
        today_emergency = []
        for j, (start, stop, energy) in enumerate(events or [(0, 0, 0.0)]):
            label = f"{clock(start)}-{clock(stop)}" if stop else "无"
            emergency.append([date if j == 0 else None, label, energy])
            record = {"date": date, "period": label, "start_slot": start if stop else None,
                      "end_slot_exclusive": stop if stop else None, "emergency_kwh": energy}
            all_emergency.append(record)
            today_emergency.append(record)
        if date not in SPECIFIED:
            continue
        summary = {"date": date, "planned_kwh": float(d["original"][i].sum()),
                   "planned_cost_yuan": float(d["fees"][i, :, 0].sum()),
                   "emergency_cost_yuan": float(d["fees"][i, :, 3].sum()),
                   "total_cost_yuan": float(d["fees"][i].sum()), "emergency_kwh": float(d["emergency"][i].sum()),
                   "initial_soc_kwh": float(d["states"][i, 0]), "final_soc_kwh": float(d["states"][i, -1])}
        summaries.append(summary)
        today_intervals = [{"date": date, "period": f"{h}:00-{h}:10", "planned_kwh": float(d["original"][i, h * 6])}
                           for h in (10, 12, 14, 16, 18, 20)]
        intervals.extend(today_intervals)
        specified_battery.extend(today_battery)
        specified_emergency.extend(today_emergency)
        markdown += [f"\n## {date}\n", "### 表 1：指定时段购电量及全天费用\n", "| 时间段 | 计划购电量 |", "|---|---:|"]
        markdown += [f"| {r['period']} | {r['planned_kwh']:.4f} |" for r in today_intervals]
        markdown += [f"\n全天计划购电量 **{summary['planned_kwh']:.4f} kWh**；全天购电费 **{summary['total_cost_yuan']:.4f} 元**，其中计划费用 {summary['planned_cost_yuan']:.4f} 元、紧急费用 {summary['emergency_cost_yuan']:.4f} 元。",
                     "\n### 表 2：储能充放电及日初、日末储电量\n", "| 时间段 | 充电量 | 放电量 |", "|---|---:|---:|"]
        markdown += [f"| {r['period']} | {r['charge_kwh']:.4f} | {r['discharge_kwh']:.4f} |" for r in today_battery]
        markdown += [f"\n0:00 储电量 {summary['initial_soc_kwh']:.4f} kWh；24:00 储电量 {summary['final_soc_kwh']:.4f} kWh。",
                     "\n### 表 3：紧急购电连续时段\n", "| 时间段 | 紧急购电量 |", "|---|---:|"]
        markdown += [f"| {r['period']} | {r['emergency_kwh']:.4f} |" for r in today_emergency]
    payload = {"dates": dates, "original": d["original"].tolist(), "final": d["final"].tolist(),
               "fees": d["fees"].sum((1, 2)).tolist(), "battery": battery, "emergency": emergency}
    write_json(WORK / "workbook-2.json", payload)
    tables = {"table1_summary": summaries, "table1_intervals": intervals,
              "table2_battery": specified_battery, "table3_emergency": specified_emergency,
              "all_battery_blocks": all_battery, "all_emergency_periods": all_emergency}
    for name, rows in tables.items():
        pd.DataFrame(rows).to_csv(table_dir / f"{name}.csv", index=False)
    write_json(table_dir / "specified_tables.json", {k: v for k, v in tables.items() if not k.startswith("all_")})
    (REPORT / "specified_dates.md").write_text("\n".join(markdown) + "\n")
    evidence = {"status": "prepared", "run_id": CASE, "role": "primary" if CASE == "regularized_42" else "cost_only_ablation", "physical_and_billing": check,
                "warmup": {**warm_check, "source": str(warm_path.relative_to(ROOT)), "sha256": sha256(warm_path),
                           "january_initial_soc_and_february_continuity": "passed"},
                "source_sha256": source_hashes,
                "dispatch_sha256": sha256(OUT / CASE / "dispatch_2.npz"),
                "daily_csv_sha256": sha256(OUT / CASE / "daily.csv"),
                "template_sha256": sha256(ROOT / "data/templates/result2.xlsx"),
                "source_actuals_and_fixed_tariff_match": "passed", "daily_cost_reconciliation": "passed",
                "specified_dates": list(SPECIFIED), "battery_rows": len(battery), "emergency_rows": len(emergency),
                "template_interval_header_correction": "B1:EO1 aligned to 00:00–00:10 through 23:50–24:00; original template is shifted by one slot",
                "efficiency_convention": "eta_charge = eta_discharge = sqrt(0.9), matching the frozen comparison protocol",
                "workbook_author": "@oai/artifact-tool via existing reports/export_workbooks_v2.mjs"}
    write_json(REPORT / "evidence/workbook_qa.json", evidence)
    return payload, evidence


def author(node, mode="export"):
    WORK.mkdir(parents=True, exist_ok=True)
    packages = Path(node).resolve().parents[1] / "node_modules"
    assert (packages / "@oai/artifact-tool").exists(), "Bundled artifact-tool required"
    link = WORK / "node_modules"
    if not link.exists():
        link.symlink_to(packages, target_is_directory=True)
    source = (ROOT / "reports/export_workbooks_v2.mjs").read_text()
    clear_emergency = '  emergency.getRange(`A2:C${Math.max(last,11)}`).clear({applyTo:"contents"});'
    assert source.count(clear_emergency) == 1
    source = source.replace(clear_emergency, clear_emergency.replace('"contents"', '"all"') + '\n  emergency.getRange(`A2:C${last}`).format.font = {name:"宋体",size:10};')
    # Template example separators mark example dates, not the exported events.
    # Keep its black column/date outlines, with boundaries at the real dates.
    needle = "  emergency.freezePanes.freezeRows(1);"
    assert source.count(needle) == 1
    date_borders = '''
  const dayStarts = input.emergency.flatMap((row,index) => row[0] ? [index+2] : []);
  for (let i=0; i<dayStarts.length; i++) {
    const start=dayStarts[i], end=i+1<dayStarts.length ? dayStarts[i+1]-1 : last;
    const line={style:"thin",color:"#000000"};
    emergency.getRange(`A${start}:C${end}`).format.borders = {
      top:line,bottom:line,left:line,right:line,insideVertical:line
    };
  }
'''
    source = source.replace(needle, needle + date_borders)
    source = source.replace('?"A1:H8":"A1:F9"', '?"A1:H8":name==="紧急购电量"?"A1:C9":"A1:F9"')
    needle = "  workbook.recalculate();"
    assert source.count(needle) == 1
    mutation = '''  const planSheet = workbook.worksheets.getItem("计划购电量");
  const mutationTests = [];
  for (const [cell, total, delta] of [["B2", "EP2", 1], ["EO168", "EP168", 2], ["EO335", "EP335", 3]]) {
    const before = planSheet.getRange(cell).values[0][0];
    const sumBefore = planSheet.getRange(total).values[0][0];
    planSheet.getRange(cell).values = [[before + delta]];
    const after = planSheet.getRange(total).values[0][0];
    if (Math.abs(after - sumBefore - delta) > 1e-6) throw new Error("Daily sum failed input mutation: " + cell);
    planSheet.getRange(cell).values = [[before]];
    mutationTests.push({cell, formulaCell:total, delta, observedDelta:after-sumBefore, restored:true});
  }
  await fs.writeFile(path.join(report,"workbook-previews","recalculation-tests.json"),JSON.stringify(mutationTests,null,2));
'''
    source = source.replace(needle, mutation + needle)
    needle = "  const output=await SpreadsheetFile.exportXlsx(workbook);"
    assert source.count(needle) == 1
    scan = r'''  const scan = await workbook.inspect({kind:"match",searchTerm:"#REF!|#DIV/0!|#VALUE!|#NAME\\?|#N/A|#NUM!|#NULL!|#SPILL!|#CALC!",options:{useRegex:true,maxResults:300},maxChars:2500,summary:"final formula error scan"});
  await fs.writeFile(path.join(report,"workbook-previews","formula-errors.json"),JSON.stringify({ndjson:scan.ndjson,records:scan.records},null,2));
'''
    source = source.replace(needle, scan + needle)
    source = source.replace('path.join(root,"data/results",runId,fileName)', 'path.join(report,fileName)')
    script = WORK / "export_workbooks.mjs"
    script.write_text(source)
    run_id = "exp007" if CASE == "regularized_42" else f"exp007/{CASE}"
    subprocess.run([str(node), str(script), str(ROOT), run_id, mode, "2"], check=True)
    generated_inspect = REPORT / "result2.xlsx.inspect.ndjson"
    if generated_inspect.exists():
        generated_inspect.replace(REPORT / "workbook-previews/saved-workbook.inspect.ndjson")


def verify(payload, evidence):
    target = REPORT / "result2.xlsx"
    saved = load_workbook(target, data_only=True, read_only=True)
    formulas = load_workbook(target, data_only=False, read_only=False)
    template = load_workbook(ROOT / "data/templates/result2.xlsx", data_only=False, read_only=False)
    assert saved.sheetnames == template.sheetnames == ["计划购电量", "充放电量", "紧急购电量"]
    plan = saved["计划购电量"]
    rows = list(plan.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 334 and all(len(row) == 147 for row in rows)
    assert [r[0].date().isoformat() for r in rows] == payload["dates"]
    np.testing.assert_allclose([r[1:145] for r in rows], payload["original"], rtol=0, atol=1e-6)
    np.testing.assert_allclose([r[145] for r in rows], np.sum(payload["original"], axis=1), rtol=0, atol=1e-6)
    np.testing.assert_allclose([r[146] for r in rows], payload["fees"], rtol=0, atol=1e-6)
    for slot in range(144):
        assert plan.cell(1, slot+2).value == f"{clock(slot)}-{clock(slot+1)}"
    for row in range(2, 336):
        assert formulas["计划购电量"].cell(row, 146).value == f"=SUM(B{row}:EO{row})"
    for name, key in (("充放电量", "battery"), ("紧急购电量", "emergency")):
        rows = list(saved[name].iter_rows(min_row=2, values_only=True))
        assert len(rows) == len(payload[key]), (name, len(rows), len(payload[key]))
        for actual_row, expected_row in zip(rows, payload[key]):
            for index, expected in enumerate(expected_row):
                if isinstance(expected, (int, float)):
                    np.testing.assert_allclose(actual_row[index], expected, rtol=0, atol=1e-6)
                elif index == 0 and expected:
                    assert actual_row[index].date().isoformat() == expected
                else:
                    assert actual_row[index] == expected
    for sheet in saved:
        assert not any(cell.data_type == "e" for row in sheet for cell in row)
    # Preserve template header labels, base fonts, fills, and border conventions.
    for sheet in template:
        authored = formulas[sheet.title]
        assert list(authored.merged_cells) == list(sheet.merged_cells)
        assert len(authored.data_validations.dataValidation) == len(sheet.data_validations.dataValidation)
        for cell in sheet[1]:
            if sheet.title == "计划购电量" and 2 <= cell.column <= 145:
                continue  # documented correction of template interval labels
            assert authored[cell.coordinate].value == cell.value
        for address in (["A1", "B1", "EP1", "EQ1"] if sheet.title == "计划购电量" else ["A1", "B1", "C1"]):
            a, b = authored[address], sheet[address]
            assert (a.font.name, a.font.sz, a.font.bold, a.fill.patternType) == (b.font.name, b.font.sz, b.font.bold, b.fill.patternType)
            for side in ("left", "right", "top", "bottom"):
                assert getattr(getattr(a.border, side), "style", None) == getattr(getattr(b.border, side), "style", None)
    sums = json.loads((REPORT / "workbook-previews/recalculation-tests.json").read_text())
    assert len(sums) == 3 and all(r["restored"] for r in sums)
    scan = json.loads((REPORT / "workbook-previews/formula-errors.json").read_text())
    assert any(r.get("message") == "Cell search matched 0 entries." for r in scan["records"])
    assert sha256(ROOT / "data/templates/result2.xlsx") == evidence["template_sha256"]
    evidence.update(status="passed", saved_workbook_readback="passed", cached_daily_sum_formulas="passed",
                    original_template_unchanged=True, template_structure_and_header_style="passed",
                    workbook_sha256=sha256(target), worksheet_names=saved.sheetnames,
                    workbook_dimensions={s.title: [s.max_row, s.max_column] for s in formulas},
                    saved_error_cells=0, artifact_formula_error_scan="0 matches", recalculation_tests=sums,
                    native_excel_recalculation="not exercised; artifact-tool input-mutation tests and saved cached values verified",
                    render_engine="artifact-tool; rendered workbook ranges, not native Excel screenshots",
                    visual_review="pending", checked_utc=datetime.now(UTC).isoformat())
    for workbook in (saved, formulas, template):
        workbook.close()
    write_json(REPORT / "evidence/workbook_qa.json", evidence)
    print(json.dumps({"status": "passed", "workbook": str(target), "total_cost_yuan": evidence["physical_and_billing"]["total_cost_yuan"],
                      "intervals": 48096, "visual_review": "pending"}, ensure_ascii=False))


def run(node, prepare_only=False, case="regularized_42", inspect_template=False):
    global CASE, REPORT, WORK
    CASE = case
    REPORT = ROOT / "reports/experiments/exp007"
    WORK = ROOT / ".work/exp007"
    if case != "regularized_42":
        REPORT = REPORT / case
        WORK = WORK / case
    began = time.monotonic()
    if inspect_template:
        author(node, mode="inspect")
        return
    payload, evidence = prepare()
    if not prepare_only:
        author(node)
        verify(payload, evidence)
        evidence["export_and_verification_seconds"] = time.monotonic() - began
        write_json(REPORT / "evidence/workbook_qa.json", evidence)
    print(f"Workbook export/check wall time: {time.monotonic()-began:.2f} s")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--node", required=True)
    parser.add_argument("--prepare-only", action="store_true")
    parser.add_argument("--case", choices=("regularized_42", "cost_only_42"), default="regularized_42")
    parser.add_argument("--inspect-template", action="store_true")
    run(**vars(parser.parse_args()))
