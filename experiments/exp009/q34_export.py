"""Export verified Q3/Q4-3 formal-period arrays into the supplied templates.

All numeric values are computed before writing; original templates contain no
formulas. The output retains that value-only convention, so Excel recalculation
is not required. The four requested sheets keep their original names.
"""
import argparse
from copy import copy
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from experiments.exp009.q34_run import OUT, ROOT, save, verify, digest
from experiments.exp009.q34_forecast import SelectedForecasts


def clock(slot):
    return f'{slot//6:02d}:{slot%6*10:02d}'


def export(scenario):
    directory=OUT/('q3' if scenario=='3' else 'q4_3')
    archive=directory/f'dispatch_{scenario}.npz'
    with np.load(archive) as z: a={key:z[key].copy() for key in z.files}
    assert np.array_equal(a['days'],np.arange(31,365))
    verified=verify(a,float(a['states'][0,0]))
    template=ROOT/f'data/templates/result{scenario}.xlsx'
    workbook=load_workbook(template)
    assert workbook.sheetnames==['计划购电量','调整购电量','充放电量','紧急购电量']
    assert not any(c.data_type=='f' for s in workbook for row in s for c in row)
    dates=pd.date_range('2025-02-01','2025-12-31')
    expected={}
    for name,key in [('计划购电量','original'),('调整购电量','final')]:
        sheet=workbook[name]
        sheet.cell(1,1,'日期\\时间')
        for slot in range(144):sheet.cell(1,slot+2,f'{clock(slot)}-{clock(slot+1)}')
        sheet.cell(1,146,'全天购电量/kWh')
        sheet.cell(1,147,'原计划购电费/元' if key=='original' else '最终总购电费/元')
        for i,date in enumerate(dates):
            fees=float(a['fees'][i,:,0].sum()) if key=='original' else float(a['fees'][i].sum())
            row=[date.to_pydatetime(),*map(float,a[key][i]),float(a[key][i].sum()),fees]
            for j,value in enumerate(row,1):
                cell=sheet.cell(i+2,j,value)
                cell.number_format='yyyy/mm/dd' if j==1 else '0.00' if j==147 else '0.0000'
        sheet.freeze_panes='B2'
        expected[name]=(334,144)
    storage=workbook['充放电量']
    styles=[copy(storage.cell(2,j)._style) for j in range(1,7)]
    for row in storage:
        if row[0].row>1:
            for c in row:c.value=None
    for i,date in enumerate(dates):
        for block in range(6):
            start,stop=block*24,(block+1)*24
            row=[date.to_pydatetime() if block==0 else None,f'{clock(start)}-{clock(stop)}',
                float(a['charge'][i,start:stop].sum()),float(a['discharge'][i,start:stop].sum()),
                '00:00' if block==0 else '24:00' if block==1 else None,
                float(a['states'][i,0]) if block==0 else float(a['states'][i,-1]) if block==1 else None]
            for j,value in enumerate(row,1):
                c=storage.cell(2+i*6+block,j,value);c._style=copy(styles[j-1])
                c.number_format='yyyy/mm/dd' if j==1 else '0.0000' if j in (3,4,6) else 'General'
    storage.freeze_panes='C2'
    emergency=workbook['紧急购电量']
    emergency_styles=[copy(emergency.cell(2,j)._style) for j in range(1,4)]
    for row in emergency:
        if row[0].row>1:
            for c in row:c.value=None
    emergency_rows=[]
    for i,date in enumerate(dates):
        mask=a['emergency'][i]>1e-6
        edges=np.diff(np.r_[False,mask,False].astype(int))
        starts,stops=np.flatnonzero(edges==1),np.flatnonzero(edges==-1)
        if not len(starts): emergency_rows.append([date.to_pydatetime(),'无',0.])
        for j,(start,stop) in enumerate(zip(starts,stops)):
            emergency_rows.append([date.to_pydatetime() if j==0 else None,
                f'{clock(int(start))}-{clock(int(stop))}',float(a['emergency'][i,start:stop].sum())])
    for i,row in enumerate(emergency_rows,2):
        for j,value in enumerate(row,1):
            c=emergency.cell(i,j,value);c._style=copy(emergency_styles[j-1])
            c.number_format='yyyy/mm/dd' if j==1 else '0.0000' if j==3 else 'General'
    output=OUT/f'result{scenario}.xlsx'
    workbook.save(output)
    saved=load_workbook(output,data_only=False)
    max_error=0.;cells=0
    for name,key in [('计划购电量','original'),('调整购电量','final')]:
        s=saved[name]
        assert s.max_row==335 and s.max_column==147
        assert s.cell(1,2).value=='00:00-00:10' and s.cell(1,145).value=='23:50-24:00'
        actual=np.array([[s.cell(i+2,j+2).value for j in range(144)] for i in range(334)])
        max_error=max(max_error,float(np.abs(actual-a[key]).max()));cells+=actual.size
        for i,date in enumerate(dates):
            assert s.cell(i+2,1).value==date.to_pydatetime()
            assert abs(s.cell(i+2,146).value-a[key][i].sum())<1e-6
            wanted=a['fees'][i,:,0].sum() if key=='original' else a['fees'][i].sum()
            assert abs(s.cell(i+2,147).value-wanted)<1e-6
    s=saved['充放电量'];battery_error=0.
    for i,date in enumerate(dates):
        assert s.cell(i*6+2,1).value==date.to_pydatetime()
        for block in range(6):
            for column,key in ((3,'charge'),(4,'discharge')):
                battery_error=max(battery_error,abs(s.cell(i*6+block+2,column).value-a[key][i,block*24:(block+1)*24].sum()))
        assert abs(s.cell(i*6+2,6).value-a['states'][i,0])<1e-6
        assert abs(s.cell(i*6+3,6).value-a['states'][i,-1])<1e-6
    for i,row in enumerate(emergency_rows,2):
        for j,wanted in enumerate(row,1):
            got=saved['紧急购电量'].cell(i,j).value
            if isinstance(wanted,float): assert abs(got-wanted)<1e-6
            else: assert got==wanted
    assert max_error<1e-6 and battery_error<1e-6
    assert not any(c.data_type=='f' for s in saved for row in s for c in row)
    save(directory/'workbook_verification.json',{'passed':True,'file':str(output.relative_to(ROOT)),
        'file_sha256':digest(output),'archive_sha256':digest(archive),'template_sha256':digest(template),
        'dates':['2025-02-01','2025-12-31'],'days':334,'purchase_numeric_cells_checked':cells,
        'max_purchase_error_kwh':max_error,'max_battery_block_error_kwh':float(battery_error),
        'emergency_rows_checked':len(emergency_rows),'formulas':0,
        'plan_fee_field':'original planned bill only','adjusted_fee_field':'total actual bill including refund and emergency',
        'all_saved_numeric_cells_and_dates_checked':True,'total_cost':verified['total_cost']})
    # Direct output for the six question-specified ten-minute intervals.
    slots=[]
    for day in (78,171,265,354):
        i=day-31
        for slot in (60,72,84,96,108,120):
            slots.append({'date':str(dates[i].date()),'interval':f'{clock(slot)}-{clock(slot+1)}',
                'original_kwh':float(a['original'][i,slot]),'final_kwh':float(a['final'][i,slot])})
    pd.DataFrame(slots).to_csv(directory/'specified_slots.csv',index=False)
    print(output,verified['total_cost'])


def forecast_metrics():
    f=SelectedForecasts(); rows=[]
    for scenario in ('3','4-3'):
        for mode in ('midnight','successive_nonoverlap_6h'):
            predictions=[]; truths=[]
            for day in range(31,365):
                if mode=='midnight':
                    p=f.get(day,0,scenario)
                    values=np.column_stack((p['load_kw'],p['pv_kw'],p['price']))
                else:
                    values=np.concatenate([np.column_stack((p['load_kw'][:36],p['pv_kw'][:36],p['price'][:36]))
                        for p in (f.get(day,slot,scenario) for slot in (0,36,72,108))])
                predictions.append(values);truths.append(f.data.actual[day*144:(day+1)*144])
            predicted,truth=np.concatenate(predictions),np.concatenate(truths)
            for channel,name in enumerate(('load','pv','price','net')):
                if name=='price' and scenario=='3':continue
                error=(predicted[:,0]-predicted[:,1])-(truth[:,0]-truth[:,1]) if name=='net' else predicted[:,channel]-truth[:,channel]
                rows.append({'scenario':scenario,'mode':mode,'target':name,'samples':len(error),
                    'unit':'yuan/kWh' if name=='price' else 'kW','mae':float(np.abs(error).mean()),
                    'rmse':float(np.sqrt(np.mean(error**2))),'bias_prediction_minus_actual':float(error.mean())})
    pd.DataFrame(rows).to_csv(OUT/'q34_forecast_metrics.csv',index=False)


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--scenario',choices=('3','4-3','both'),default='both')
    p.add_argument('--forecast-metrics',action='store_true');args=p.parse_args()
    for scenario in (('3','4-3') if args.scenario=='both' else (args.scenario,)):export(scenario)
    if args.forecast_metrics:forecast_metrics()
