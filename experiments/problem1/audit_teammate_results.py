"""Read-only numerical audit of teammate workbooks; never saves an XLSX."""
from pathlib import Path
from openpyxl import load_workbook
import math,json
ROOT=Path(__file__).resolve().parents[2]
P=ROOT/'目前整体思路最新-邱志烨'
eta=math.sqrt(.9); dt=1/6
def runs(z, wanted=True):
    n=len(z); starts=[i for i in range(n) if z[i]==wanted and z[(i-1)%n]!=wanted]
    if not starts:return [n] if all(x==wanted for x in z) else []
    return [next((k for k in range(1,n+1) if z[(i+k)%n]!=wanted),n) for i in starts]
def extract(name):
    w=load_workbook(P/name,data_only=True)
    rr=list(w['计算明细'].values)[1:145]
    p,L,V,g,c,d,waste,E0,E1,G,F=map(list,zip(*[r[1:12] for r in rr]))
    net=[a-b for a,b in zip(c,d)]
    tv=sum(abs(net[i]-net[i-1]) for i in range(144))
    zs=[[v>1e-5 for v in powers] for powers in [c,d]]
    on=sum((runs(z) for z in zs),[]); off=sum((runs(z,False) for z in zs),[])
    state=max(abs(E1[i]-E0[i]-eta*c[i]*dt+d[i]*dt/eta) for i in range(144))
    bal=max(abs(g[i]+V[i]+d[i]-L[i]-c[i]-waste[i]) for i in range(144))
    charge=sum(c)*dt; discharge=sum(d)*dt; loss=charge-discharge
    cost=sum(p[i]*g[i]*dt for i in range(144))
    summary=dict(file=name,cost=cost,purchase=sum(g)*dt,charge=charge,discharge=discharge,loss=loss,
      load=sum(L)*dt,pv=sum(V)*dt,waste=sum(waste)*dt,tv=tv,starts=len(on),
      short_runs=sum(k<3 for k in on),short_deficit=sum(max(0,3-k) for k in on),
      min_on=min(on),min_off=min(off),low=sum(1e-5<abs(x)<50 for x in net),
      max_step=max(abs(net[i]-net[i-1]) for i in range(144)),
      balance=bal,state=state,continuity=max(abs(E1[i]-E0[i+1]) for i in range(143)),
      end=max(abs(E0[0]-6000),abs(E1[-1]-6000)),
      Emin=min(E0+E1),Emax=max(E0+E1),max_c=max(c),max_d=max(d),
      simultaneous=max(min(a,b) for a,b in zip(c,d)),
      daily_identity=sum(g)*dt-(sum(L)-sum(V))*dt-loss-sum(waste)*dt,
      power_min_active=min(abs(x) for x in net if abs(x)>1e-5),
      specified={rr[i][0]:g[i]*dt for i in [60,72,84,96,108,120]},
      blocks=[[j*4,(j+1)*4,sum(c[j*24:(j+1)*24])*dt,sum(d[j*24:(j+1)*24])*dt] for j in range(6)],
      no_storage_cost=sum(p[i]*max(L[i]-V[i],0)*dt for i in range(144)),
      no_storage_purchase=sum(max(L[i]-V[i],0)*dt for i in range(144)))
    summary['plan_error']=max(abs(w['计划购电量'].cell(i+2,2).value-g[i]*dt) for i in range(144))
    summary['block_error']=max(abs(w['充放电量'].cell(j+2,k+2).value-summary['blocks'][j][k+2]) for j in range(6) for k in range(2))
    return summary,rr
old,_=extract('result1.xlsx'); new,rr=extract('result1_修正版.xlsx')
w=load_workbook(P/'result1_敏感度分析.xlsx',data_only=True)
sensitivity=[]
for i,r in enumerate(list(w['本轮结果汇总'].values)[1:]):
    col=2 if i==0 else i+3
    amounts=[w['计划购电量'].cell(t+2,col).value for t in range(144)]
    cost=sum(amounts[t]*rr[t][1] for t in range(144))
    sensitivity.append(dict(id=r[0],purchase=sum(amounts),cost=cost,total_error=sum(amounts)-r[1],
                           starts=r[6],tv=r[7],step=r[8],short=r[9],low=r[10]))
raw=list(load_workbook(ROOT/'data/raw/附件1.xlsx',data_only=True).active.values)[1:145]
input_error=max(abs(float(raw[i][j+1])-rr[i][j+1]) for i in range(144) for j in range(3))
out=dict(original=old,revised=new,sensitivity=sensitivity,input_error=input_error,
 tv_reduction=(1-new['tv']/old['tv'])*100,cost_increase=new['cost']-old['cost'],
 cost_increase_pct=(new['cost']/old['cost']-1)*100,
 no_storage_save_pct=(1-new['cost']/new['no_storage_cost'])*100)
print(json.dumps(out,ensure_ascii=False,indent=2))

