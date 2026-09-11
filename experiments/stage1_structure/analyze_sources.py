from __future__ import annotations

import json
from pathlib import Path
from statistics import mean

from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[2]
PDF = ROOT / "C题.pdf"
RAW = ROOT / "data" / "raw"
TEMPLATES = ROOT / "data" / "templates"


def fmt(value):
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def workbook_summary(path: Path) -> dict:
    wb = load_workbook(path, read_only=False, data_only=False)
    result = {"file": str(path.relative_to(ROOT)), "sheets": []}
    for ws in wb.worksheets:
        nonempty = []
        formulas = 0
        numeric = []
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                nonempty.append((cell.coordinate, fmt(cell.value)))
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas += 1
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    numeric.append(float(cell.value))
        sheet = {
            "name": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "nonempty_cells": len(nonempty),
            "formula_cells": formulas,
            "merged_ranges": [str(r) for r in ws.merged_cells.ranges],
            "first_30_nonempty": nonempty[:30],
            "last_12_nonempty": nonempty[-12:],
        }
        if numeric:
            sheet["numeric_summary"] = {
                "count": len(numeric),
                "min": min(numeric),
                "max": max(numeric),
                "mean": mean(numeric),
            }
        result["sheets"].append(sheet)
    return result


def pdf_summary() -> dict:
    reader = PdfReader(PDF)
    return {
        "file": str(PDF.relative_to(ROOT)),
        "pages": len(reader.pages),
        "page_text": [page.extract_text() or "" for page in reader.pages],
    }


def main():
    payload = {
        "pdf": pdf_summary(),
        "workbooks": [
            workbook_summary(path)
            for path in sorted(list(RAW.glob("*.xlsx")) + list(TEMPLATES.glob("*.xlsx")))
        ],
    }
    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
