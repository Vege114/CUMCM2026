from __future__ import annotations

from pathlib import Path
from collections import Counter
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]


def show_cell(ws, coord):
    value = ws[coord].value
    if hasattr(value, "isoformat"):
        value = value.isoformat()
    return f"{coord}={value!r}"


def main():
    paths = sorted((ROOT / "data" / "raw").glob("*.xlsx"))
    for path in paths:
        print(f"\n[{path.name}]")
        wb = load_workbook(path, read_only=False, data_only=True)
        for ws in wb.worksheets:
            print(f"sheet={ws.title!r} shape={ws.max_row}x{ws.max_column}")
            coords = ["A1", "B1", "C1", "D1", "A2", "B2"]
            if ws.max_column >= 26:
                coords += [f"{ws.cell(1, ws.max_column).coordinate}", f"{ws.cell(2, ws.max_column).coordinate}"]
            coords += [f"A{ws.max_row}", f"{ws.cell(ws.max_row, ws.max_column).coordinate}"]
            print("  " + "; ".join(show_cell(ws, c) for c in coords))
            blanks = 0
            types = Counter()
            numeric = []
            for row in ws.iter_rows(min_row=2, min_col=2):
                for cell in row:
                    value = cell.value
                    if value is None:
                        blanks += 1
                    else:
                        types[type(value).__name__] += 1
                        if isinstance(value, (int, float)) and not isinstance(value, bool):
                            numeric.append(float(value))
            print(f"  data_blanks={blanks} types={dict(types)}")
            if numeric:
                print(f"  numeric_count={len(numeric)} min={min(numeric):.6g} max={max(numeric):.6g}")

    print("\n[templates]")
    for path in sorted((ROOT / "data" / "templates").glob("*.xlsx")):
        wb = load_workbook(path, read_only=False, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"{ws.title}:{ws.max_row}x{ws.max_column}")
        print(f"{path.name}: " + ", ".join(parts))


if __name__ == "__main__":
    main()
