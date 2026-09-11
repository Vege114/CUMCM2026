"""Populate copies of the four competition templates and verify saved values."""

import argparse
import copy
import hashlib
import json

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from .data import EPOCH, ROOT
from .dispatch import ETA, MAX_SOC, MIN_SOC, POWER_ENERGY


def clock_label(slot):
    minutes = int(slot) * 10
    return f"{minutes // 60:02d}:{minutes % 60:02d}"


def emergency_periods(energy):
    active = np.asarray(energy) > 1e-6
    edges = np.diff(np.r_[False, active, False].astype(int))
    return [(int(a), int(b), float(np.sum(energy[a:b])))
            for a, b in zip(np.where(edges == 1)[0], np.where(edges == -1)[0])]


def independent_check(d):
    for name in ("original", "final", "charge", "discharge", "emergency", "surplus"):
        assert d[name].shape == (334, 144)
        assert np.isfinite(d[name]).all() and np.min(d[name]) >= -1e-7
    assert d["states"].shape == (334, 145)
    np.testing.assert_allclose(d['states'][:-1, -1], d['states'][1:, 0], atol=1e-6)
    assert d['states'].min() >= MIN_SOC - 1e-6
    assert d['states'].max() <= MAX_SOC + 1e-6
    assert max(d['charge'].max(), d['discharge'].max()) <= POWER_ENERGY + 1e-6
    assert not np.any((d['charge'] > 1e-6) & (d['discharge'] > 1e-6))
    np.testing.assert_allclose(np.diff(d['states'], axis=1),
                               ETA * d['charge'] - d['discharge'] / ETA, atol=1e-6)
    balance = (d['final'] + d['actual'][:, :, 1] / 6 + d['discharge'] + d['emergency']
               - d['actual'][:, :, 0] / 6 - d['charge'] - d['surplus'])
    assert np.abs(balance).max() < 1e-6
    delta = d['final'] - d['original']
    expected = np.stack((d['original'] * d['price'],
                         np.maximum(delta, 0) * 1.5 * d['price'],
                         np.maximum(-delta, 0) * .5 * d['price'],
                         d['emergency'] * 5 * d['price']), axis=-1)
    np.testing.assert_allclose(expected, d['fees'], atol=1e-6)
    return {"days": 334, "intervals_per_day": 144, "violations": 0,
            "max_energy_residual": float(np.abs(balance).max()),
            "cost_independent_recalculation": "passed",
            "total_cost": float(expected.sum())}


def main():
    p = argparse.ArgumentParser(); p.add_argument("--run-id", default="exp001")
    args = p.parse_args()
    out = ROOT / "data/results" / args.run_id
    report = ROOT / "reports/experiments" / args.run_id
    report.mkdir(parents=True, exist_ok=True)
    verification, table_rows, emergency_rows, battery_rows = {}, [], [], []
    specified = {"2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"}
    appendix = ["# 指定日期完整结果\n", "数值保留四位小数；精确值保存在结果工作簿和 CSV。购电量及储电量单位为千瓦时，费用单位为元。\n"]
    for scenario in ("2", "3", "4-2", "4-3"):
        file_name = f"result{scenario}.xlsx"
        with np.load(out / f"dispatch_{scenario}.npz") as a:
            d = {k: a[k] for k in a.files}
        check = independent_check(d)
        template_path = ROOT / "data/templates" / file_name
        template_hash = hashlib.sha256(template_path.read_bytes()).hexdigest()
        w = load_workbook(template_path)
        for sheet_name, key in (("计划购电量", "original"), ("调整购电量", "final")):
            if sheet_name not in w.sheetnames:
                continue
            s = w[sheet_name]
            for slot in range(144):
                s.cell(1, slot + 2, f"{clock_label(slot)}-{clock_label(slot + 1)}")
            for day in range(334):
                row = day + 2
                date = EPOCH + pd.Timedelta(days=day + 31)
                s.cell(row, 1, date.to_pydatetime()).number_format = 'yyyy/mm/dd'
                for slot, value in enumerate(d[key][day]):
                    s.cell(row, slot + 2, float(value)).number_format = '0.0000'
                s.cell(row, 146, float(d[key][day].sum())).number_format = '0.0000'
                s.cell(row, 147, float(d['fees'][day].sum())).number_format = '0.00'
            s.freeze_panes = "B2"
            s.column_dimensions['A'].width = 14
            for col in range(2, 148):
                s.column_dimensions[get_column_letter(col)].width = 18
            s.row_dimensions[1].height = 32
            for cell in s[1]:
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        s = w['充放电量']
        styles = [[copy.copy(c._style) for c in row] for row in s.iter_rows(min_row=2,max_row=7)]
        s.delete_rows(2, s.max_row)
        for day in range(334):
            date = EPOCH + pd.Timedelta(days=day + 31)
            for block in range(6):
                row = 2 + day * 6 + block
                values = [date.to_pydatetime() if block == 0 else None,
                          f"{block*4}:00-{(block+1)*4}:00",
                          float(d['charge'][day, block*24:(block+1)*24].sum()),
                          float(d['discharge'][day, block*24:(block+1)*24].sum()),
                          "0:00" if block == 0 else "24:00" if block == 1 else None,
                          float(d['states'][day, 0]) if block == 0 else
                          float(d['states'][day, -1]) if block == 1 else None]
                for col, value in enumerate(values, 1):
                    cell = s.cell(row, col, value); cell._style = copy.copy(styles[block][col-1])
                    if col in (3,4,6): cell.number_format = '0.0000'
                    if col == 1: cell.number_format = 'yyyy/mm/dd'
                battery_rows.append({"scenario": scenario, "date": str(date.date()),
                                     "period": values[1], "charge_kwh": values[2],
                                     "discharge_kwh": values[3],
                                     "initial_soc": float(d['states'][day,0]),
                                     "final_soc": float(d['states'][day,-1])})
        s.freeze_panes = "C2"
        for col, width in zip('ABCDEF', (14,20,18,18,12,18)): s.column_dimensions[col].width=width
        s = w['紧急购电量']; style = [copy.copy(c._style) for c in s[2]]
        s.delete_rows(2,s.max_row); row = 2
        for day in range(334):
            date = EPOCH + pd.Timedelta(days=day+31)
            periods = emergency_periods(d['emergency'][day])
            np.testing.assert_allclose(sum(p[2] for p in periods), d['emergency'][day].sum(), atol=1e-5)
            for i,(start,stop,energy) in enumerate(periods or [(0,0,0.0)]):
                label = f"{clock_label(start)}-{clock_label(stop)}" if stop else "无"
                for col,value in enumerate((date.to_pydatetime() if i==0 else None,label,energy),1):
                    cell=s.cell(row,col,value); cell._style=copy.copy(style[col-1])
                    cell.number_format='yyyy/mm/dd' if col==1 else '0.0000' if col==3 else 'General'
                emergency_rows.append({"scenario":scenario,"date":str(date.date()),
                                       "period":label,"emergency_kwh":energy})
                row+=1
            if str(date.date()) in specified:
                values={"scenario":scenario,"date":str(date.date()),
                        "planned_kwh":float(d['original'][day].sum()),
                        "final_kwh":float(d['final'][day].sum()),
                        "total_cost":float(d['fees'][day].sum()),
                        "emergency_kwh":float(d['emergency'][day].sum()),
                        "initial_soc":float(d['states'][day,0]),"final_soc":float(d['states'][day,-1])}
                for hour in (10,12,14,16,18,20):
                    values[f"grid_{hour:02d}"]=float(d['final'][day,hour*6])
                    values[f"plan_{hour:02d}"]=float(d['original'][day,hour*6])
                table_rows.append(values)
                appendix += [f"\n## 问题 {scenario} · {date.date()}\n",
                             "### 表 1：指定时段购电与全天结果\n",
                             "| 时间段 | 凌晨计划购电量 | 最终购电量 |\n|---|---:|---:|"]
                appendix += [f"| {hour:02d}:00-{hour:02d}:10 | {values[f'plan_{hour:02d}']:.4f} | {values[f'grid_{hour:02d}']:.4f} |" for hour in (10,12,14,16,18,20)]
                appendix += [f"\n全天计划购电量：{values['planned_kwh']:.4f}；最终购电量：{values['final_kwh']:.4f}；总购电费：{values['total_cost']:.4f}。\n",
                             "### 表 2：储能充放电与日初日末储电量\n",
                             "| 时间段 | 充电量 | 放电量 |\n|---|---:|---:|"]
                appendix += [f"| {r['period']} | {r['charge_kwh']:.4f} | {r['discharge_kwh']:.4f} |"
                             for r in battery_rows if r['scenario'] == scenario
                             and r['date'] == str(date.date())]
                appendix += [f"\n0:00 储电量：{values['initial_soc']:.4f}；24:00 储电量：{values['final_soc']:.4f}。\n",
                             "### 表 3：紧急购电明细\n", "| 时间段 | 紧急购电量 |\n|---|---:|"]
                appendix += [f"| {clock_label(a)}-{clock_label(b)} | {e:.4f} |" for a,b,e in periods] or ["| 无 | 0.0000 |"]
        s.freeze_panes='C2'
        for col,width in zip('ABC',(14,24,20)):s.column_dimensions[col].width=width
        w.save(out/file_name)
        assert hashlib.sha256(template_path.read_bytes()).hexdigest()==template_hash
        saved=load_workbook(out/file_name,read_only=True,data_only=True)
        for title,key in (("计划购电量","original"),("调整购电量","final")):
            if title not in saved.sheetnames:continue
            sheet=saved[title]
            rows=list(sheet.iter_rows(min_row=2,values_only=True))
            assert len(rows)==334
            np.testing.assert_allclose(np.array([r[1:145] for r in rows],float),d[key],atol=1e-6)
            np.testing.assert_allclose([r[145] for r in rows],d[key].sum(1),atol=1e-6)
            np.testing.assert_allclose([r[146] for r in rows],d['fees'].sum((1,2)),atol=1e-6)
            assert sheet.cell(1,2).value=='00:00-00:10'
            assert sheet.cell(1,145).value=='23:50-24:00'
            assert [r[0].date().isoformat() for r in rows] == [
                (EPOCH + pd.Timedelta(days=j + 31)).date().isoformat() for j in range(334)]
        battery = list(saved['充放电量'].iter_rows(min_row=2, values_only=True))
        assert len(battery) == 334 * 6
        np.testing.assert_allclose([r[2] for r in battery], d['charge'].reshape(334 * 6, 24).sum(1), atol=1e-6)
        np.testing.assert_allclose([r[3] for r in battery], d['discharge'].reshape(334 * 6, 24).sum(1), atol=1e-6)
        np.testing.assert_allclose([battery[j * 6][5] for j in range(334)], d['states'][:, 0], atol=1e-6)
        np.testing.assert_allclose([battery[j * 6 + 1][5] for j in range(334)], d['states'][:, -1], atol=1e-6)
        emergency = list(saved['紧急购电量'].iter_rows(min_row=2, values_only=True))
        expected_emergency = [r for r in emergency_rows if r['scenario'] == scenario]
        assert len(emergency) == len(expected_emergency)
        current_date = None
        for actual_row, expected_row in zip(emergency, expected_emergency):
            if actual_row[0] is not None:
                current_date = actual_row[0].date().isoformat()
            assert current_date == expected_row['date'] and actual_row[1] == expected_row['period']
            np.testing.assert_allclose(actual_row[2], expected_row['emergency_kwh'], atol=1e-6)
        saved.close()
        check['saved_workbook_readback']='passed';check['original_template_unchanged']=True
        check['sha256']=hashlib.sha256((out/file_name).read_bytes()).hexdigest()
        verification[file_name]=check
        print('EXPORTED',file_name,flush=True)
    pd.DataFrame(table_rows).to_csv(out/'specified_dates.csv',index=False)
    pd.DataFrame(battery_rows).to_csv(out/'battery_blocks.csv',index=False)
    pd.DataFrame(emergency_rows).to_csv(out/'emergency_periods.csv',index=False)
    (report/'specified_dates.md').write_text('\n'.join(appendix))
    (out/'verification.json').write_text(json.dumps(verification,ensure_ascii=False,indent=2))
    (out/'README.md').write_text('''# 首轮正式答题结果

四个工作簿均由同一组按月验证选模、三个随机种子平均预测的调度回放生成。

- 数据时标为区间终点。导出表头统一为 00:00—00:10 至 23:50—24:00，原始模板未修改。
- “计划购电量”是凌晨原计划；“调整购电量”是各时段执行前最后一次调整后的最终电量，不是增减量。
- 购电表的“全天购电量”汇总该工作表的电量；“全天购电费”均记录包含计划、调整和紧急购电的全天总费用。多个工作表的费用不可再相加。
- “充放电量”记录实际执行，每天六个四小时时段；储电量连续跨日传递。
- 紧急购电连续时段合并，零事件日期记为“无”。
- dispatch_*.npz 保存完整逐时段计划、最终购电、执行、储电量与费用分项，daily_metrics.csv 保存逐日指标。
- ensemble_predictions.npz 与 prediction_archive.json 共同定义逐发布时刻、目标区间、模型和目标变量的全部平均预测。
- 验证结果见 verification.json；指定日期的论文表格见报告的 specified_dates.md。
''')


if __name__=='__main__':main()
