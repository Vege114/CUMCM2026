"""Prepare audited result tables; author via artifact-tool; independently read back workbooks."""

import argparse
import hashlib
import json
import shutil
import subprocess
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .data import EPOCH, ROOT
from .physics import ETA, MAX_SOC, MIN_SOC, POWER_ENERGY

SCENARIOS = ("2", "3", "4-2", "4-3")
SPECIFIED = ("2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21")


def clock(slot):
    return f"{slot//6:02d}:{slot%6*10:02d}"


def periods(energy):
    edge = np.diff(np.r_[False, np.asarray(energy) > 1e-6, False].astype(int))
    return [(int(a), int(b), float(np.sum(energy[a:b])))
            for a, b in zip(np.flatnonzero(edge == 1), np.flatnonzero(edge == -1))]


def read_npz(path):
    with np.load(path) as file:
        return {k: file[k].copy() for k in file.files}


def independent_check(d, days=334):
    for name in ("original", "final", "charge", "discharge", "emergency", "surplus"):
        assert d[name].shape == (days, 144)
        assert np.isfinite(d[name]).all() and d[name].min() >= -1e-6
    assert d["states"].shape == (days, 145)
    np.testing.assert_allclose(d["states"][:-1, -1], d["states"][1:, 0], atol=1e-6)
    assert d["states"].min() >= MIN_SOC - 1e-6
    assert d["states"].max() <= MAX_SOC + 1e-6
    assert max(d["charge"].max(), d["discharge"].max()) <= POWER_ENERGY + 1e-6
    assert not ((d["charge"] > 1e-6) & (d["discharge"] > 1e-6)).any()
    np.testing.assert_allclose(np.diff(d["states"], axis=1),
                               ETA*d["charge"] - d["discharge"]/ETA, atol=1e-6)
    balance = (d["final"] + d["actual"][:, :, 1]/6 + d["discharge"] + d["emergency"]
               - d["actual"][:, :, 0]/6 - d["charge"] - d["surplus"])
    assert np.abs(balance).max() < 1e-6
    # Independent vectorized greedy check, not a call back into the replay controller.
    net = d["final"] + (d["actual"][:, :, 1] - d["actual"][:, :, 0])/6
    expected_c = np.minimum(np.maximum(net, 0), np.minimum(POWER_ENERGY, (MAX_SOC-d["states"][:, :-1])/ETA))
    expected_d = np.minimum(np.maximum(-net, 0), np.minimum(POWER_ENERGY, (d["states"][:, :-1]-MIN_SOC)*ETA))
    np.testing.assert_allclose(d["charge"], expected_c, atol=1e-6)
    np.testing.assert_allclose(d["discharge"], expected_d, atol=1e-6)
    delta = d["final"] - d["original"]
    fees = np.stack((d["original"]*d["price"], 1.5*np.maximum(delta, 0)*d["price"],
                     .5*np.maximum(-delta, 0)*d["price"], 5*d["emergency"]*d["price"]), axis=-1)
    np.testing.assert_allclose(d["fees"], fees, atol=1e-6)
    return {"days": days, "intervals_per_day": 144, "violations": 0,
            "max_energy_residual": float(np.abs(balance).max()),
            "causal_greedy_controller": "passed", "total_cost": float(fees.sum())}


def prepare(run_id="exp002", scenarios=SCENARIOS):
    out = ROOT / "data/results" / run_id
    report = ROOT / "reports/experiments" / run_id
    work = ROOT / ".work" / run_id
    work.mkdir(parents=True, exist_ok=True)
    verification, specified, interval_rows, battery_rows, emergency_rows = {}, [], [], [], []
    markdown = ["# 指定四日完整结果", "电量与储电量单位为千瓦时，费用为元。精确值保留在 CSV 和工作簿中。"]
    for scenario in scenarios:
        d = read_npz(out / f"dispatch_{scenario}.npz")
        warm = read_npz(out / f"warmup_{scenario}.npz")
        check = independent_check(d)
        independent_check(warm, 31)
        np.testing.assert_allclose(warm["states"][0, 0], 6000)
        np.testing.assert_allclose(warm["states"][-1, -1], d["states"][0, 0], atol=1e-6)
        dates = [(EPOCH + pd.Timedelta(days=i+31)).date().isoformat() for i in range(334)]
        battery, emergency = [], []
        for i, date in enumerate(dates):
            for block in range(6):
                c = float(d["charge"][i, block*24:(block+1)*24].sum())
                discharge = float(d["discharge"][i, block*24:(block+1)*24].sum())
                period = f"{block*4}:00-{(block+1)*4}:00"
                battery.append([date if block == 0 else None, period, c, discharge,
                                "0:00" if block == 0 else "24:00" if block == 1 else None,
                                float(d["states"][i, 0]) if block == 0 else float(d["states"][i, -1]) if block == 1 else None])
                battery_rows.append({"scenario": scenario, "date": date, "period": period, "charge_kwh": c,
                                         "discharge_kwh": discharge, "initial_soc": float(d["states"][i, 0]),
                                         "final_soc": float(d["states"][i, -1])})
            events = periods(d["emergency"][i])
            np.testing.assert_allclose(sum(e[2] for e in events), d["emergency"][i].sum(), atol=1e-5)
            for j, (start, stop, energy) in enumerate(events or [(0, 0, 0.)]):
                label = f"{clock(start)}-{clock(stop)}" if stop else "无"
                emergency.append([date if j == 0 else None, label, energy])
                emergency_rows.append({"scenario": scenario, "date": date, "period": label, "emergency_kwh": energy})
            if date in SPECIFIED:
                summary = {"scenario": scenario, "date": date, "planned_kwh": float(d["original"][i].sum()),
                               "final_kwh": float(d["final"][i].sum()), "total_cost": float(d["fees"][i].sum()),
                               "emergency_kwh": float(d["emergency"][i].sum()), "initial_soc": float(d["states"][i, 0]),
                               "final_soc": float(d["states"][i, -1])}
                selected = []
                for hour in (10, 12, 14, 16, 18, 20):
                    original, final = float(d["original"][i, hour*6]), float(d["final"][i, hour*6])
                    summary[f"plan_{hour:02d}"], summary[f"grid_{hour:02d}"] = original, final
                    selected.append({"scenario": scenario, "date": date, "period": f"{hour}:00-{hour}:10", "original": original, "final": final})
                specified.append(summary)
                interval_rows.extend(selected)
                markdown += [f"\n## 问题 {scenario} · {date}\n", "### 表 1：购电与全天费用\n", "| 时段 | 凌晨原计划 | 最终购电 |", "|---|---:|---:|"]
                markdown += [f"| {r['period']} | {r['original']:.4f} | {r['final']:.4f} |" for r in selected]
                markdown += [f"\n全天原计划 {summary['planned_kwh']:.4f}；最终购电 {summary['final_kwh']:.4f}；全部费用 {summary['total_cost']:.4f}。",
                             "\n### 表 2：储能充放电\n", "| 时段 | 充电量 | 放电量 |", "|---|---:|---:|"]
                markdown += [f"| {r['period']} | {r['charge_kwh']:.4f} | {r['discharge_kwh']:.4f} |" for r in battery_rows[-6:]]
                markdown += [f"\n日初储电量 {summary['initial_soc']:.4f}；日末储电量 {summary['final_soc']:.4f}。",
                             "\n### 表 3：紧急购电\n", "| 时段 | 紧急购电量 |", "|---|---:|"]
                markdown += [f"| {clock(a)}-{clock(b)} | {e:.4f} |" for a, b, e in events] or ["| 无 | 0.0000 |"]
        payload = {"dates": dates, "original": d["original"].tolist(), "final": d["final"].tolist(),
                       "fees": d["fees"].sum((1, 2)).tolist(), "battery": battery, "emergency": emergency}
        (work / f"workbook-{scenario}.json").write_text(json.dumps(payload, ensure_ascii=False, allow_nan=False))
        check["january_warmup_and_february_continuity"] = "passed"
        check["template_sha256"] = hashlib.sha256((ROOT / "data/templates" / f"result{scenario}.xlsx").read_bytes()).hexdigest()
        verification[f"result{scenario}.xlsx"] = check
    for name, rows in (("specified_dates", specified), ("specified_intervals", interval_rows),
                       ("battery_blocks", battery_rows), ("emergency_periods", emergency_rows)):
        pd.DataFrame(rows).to_csv(out / f"{name}.csv", index=False)
    (report / "specified_dates.md").write_text("\n".join(markdown))
    (out / "verification.json").write_text(json.dumps(verification, indent=2))


def verify(run_id="exp002", scenarios=SCENARIOS):
    out = ROOT / "data/results" / run_id
    results = json.loads((out / "verification.json").read_text())
    for scenario in scenarios:
        file_name = f"result{scenario}.xlsx"
        payload = json.loads((ROOT / ".work" / run_id / f"workbook-{scenario}.json").read_text())
        saved = load_workbook(out / file_name, read_only=True, data_only=True)
        for name, key in (("计划购电量", "original"), ("调整购电量", "final")):
            if name not in saved.sheetnames:
                continue
            sheet = saved[name]
            rows = list(sheet.iter_rows(min_row=2, values_only=True))
            assert len(rows) == 334
            assert [r[0].date().isoformat() for r in rows] == payload["dates"]
            np.testing.assert_allclose([r[1:145] for r in rows], payload[key], atol=1e-6)
            np.testing.assert_allclose([r[145] for r in rows], np.sum(payload[key], axis=1), atol=1e-6)
            np.testing.assert_allclose([r[146] for r in rows], payload["fees"], atol=1e-6)
            assert sheet.cell(1, 2).value == "00:00-00:10"
            assert sheet.cell(1, 145).value == "23:50-24:00"
        for name, key in (("充放电量", "battery"), ("紧急购电量", "emergency")):
            rows = list(saved[name].iter_rows(min_row=2, values_only=True))
            assert len(rows) == len(payload[key])
            for a, b in zip(rows, payload[key]):
                for index, value in enumerate(b):
                    if isinstance(value, (int, float)):
                        np.testing.assert_allclose(a[index], value, atol=1e-6)
                    elif index == 0 and value:
                        assert a[index].date().isoformat() == value
                    else:
                        assert a[index] == value
        for sheet in saved:
            assert not any(c.data_type == "e" for row in sheet for c in row)
        saved.close()
        assert hashlib.sha256((ROOT / "data/templates" / file_name).read_bytes()).hexdigest() == results[file_name]["template_sha256"]
        results[file_name].update(saved_workbook_readback="passed", original_template_unchanged=True,
                                  sha256=hashlib.sha256((out / file_name).read_bytes()).hexdigest())
    (out / "verification.json").write_text(json.dumps(results, indent=2))
    print(f"Verified {len(scenarios)} workbooks, 334 days each, all physical and billing identities")


def run(run_id="exp002", node=None):
    prepare(run_id)
    if not node:
        raise RuntimeError("Workbook export requires --node with the bundled Codex Node executable")
    runtime = Path(node).resolve().parents[1] / "node_modules"
    work = ROOT / ".work" / run_id
    link = work / "node_modules"
    if not link.exists():
        link.symlink_to(runtime, target_is_directory=True)
    script = work / "export_workbooks_v2.mjs"
    shutil.copy2(ROOT / "reports/export_workbooks_v2.mjs", script)
    subprocess.run([node, str(script), str(ROOT), run_id], check=True)
    verify(run_id)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--run-id", default="exp002")
    parser.add_argument("--node")
    parser.add_argument("--prepare-only", action="store_true")
    parser.add_argument("--verify-only", action="store_true")
    parser.add_argument("--scenarios", nargs="+", choices=SCENARIOS, default=list(SCENARIOS))
    args = parser.parse_args()
    if args.prepare_only:
        prepare(args.run_id, args.scenarios)
    elif args.verify_only:
        verify(args.run_id, args.scenarios)
    else:
        run(args.run_id, args.node)


if __name__ == "__main__":
    main()
