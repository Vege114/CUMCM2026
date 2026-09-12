"""Build a read-only Q2 discussion snapshot and independently audit its numbers.

Run from any directory with the project Python environment. No experiment is
trained, replayed, selected, or changed. JSON contains only the supplied data,
the stated purchase-cost objective, and outputs explicitly requested in Q2.
"""

import argparse
import csv
import hashlib
import json
from datetime import date, timedelta
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
EPOCH = date(2025, 1, 1)
DATES = [(EPOCH + timedelta(days=i)).isoformat() for i in range(365)]
FORMAL = DATES[31:]
SPECIFIED = ["2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"]
ADDITIVE = ("total_cost", "planned_cost", "emergency_cost", "planned_kwh",
            "emergency_kwh", "charge_kwh", "discharge_kwh")
CASE_FILES = {
    "v2_primary": "data/results/exp002/dispatch_2.npz",
    "exp003_uncalibrated": "data/results/exp003/dispatch_uncalibrated_seed_42.npz",
    "exp003_primary": "data/results/exp003/dispatch_primary_seed_42.npz",
    "periodic": "data/results/exp003/dispatch_periodic_seed_42.npz",
}
CASE_LABELS = {
    "v2_primary": "v2 正式方案",
    "exp003_uncalibrated": "exp003 未校准网络（α=1,1）",
    "exp003_primary": "exp003 当前方案（α=0,0）",
    "periodic": "周期预测基线",
}


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def clock(slot):
    return f"{slot // 6:02d}:{slot % 6 * 10:02d}"


def close(actual, expected, atol=1e-6):
    np.testing.assert_allclose(actual, expected, rtol=0, atol=atol)


def csv_rows(relative):
    with (ROOT / relative).open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def read_raw():
    arrays = []
    for name in ("附件2_小区负载.csv", "附件2_光伏发电实际功率.csv"):
        rows = csv_rows(f"data/raw/{name}")
        keys = list(rows[0])
        assert [row[keys[0]] for row in rows] == DATES
        assert len(keys) == 145
        # CSV time labels are ends; the final 0:00+1 is the day's 24:00.
        def minutes(label):
            if label == "0:00+1":
                return 1440
            hour, minute = map(int, label.split(":")[:2])
            return hour * 60 + minute
        assert list(map(minutes, keys[1:])) == list(range(10, 1441, 10))
        arrays.append(np.asarray([[float(row[k]) for k in keys[1:]] for row in rows]))
    reference = csv_rows("data/raw/附件1.csv")
    assert len(reference) == 144
    price = np.asarray([float(row["电价"]) for row in reference])
    for array in (*arrays, price):
        assert np.isfinite(array).all() and array.min() >= 0
    return arrays[0], arrays[1], price


def load_dispatch(relative):
    with np.load(ROOT / relative, allow_pickle=False) as archive:
        return {key: archive[key].copy() for key in archive.files}


def summary(d, index):
    indices = np.atleast_1d(index)
    fees = d["fees"][indices].sum(axis=(0, 1))
    return {
        "total_cost": float(fees.sum()), "planned_cost": float(fees[0]),
        "emergency_cost": float(fees[3]),
        "planned_kwh": float(d["original"][indices].sum()),
        "emergency_kwh": float(d["emergency"][indices].sum()),
        "charge_kwh": float(d["charge"][indices].sum()),
        "discharge_kwh": float(d["discharge"][indices].sum()),
        "initial_soc": float(d["states"][indices[0], 0]),
        "final_soc": float(d["states"][indices[-1], -1]),
    }


def emergencies(energy):
    rows = []
    start = None
    for i in range(145):
        active = i < 144 and energy[i] > 1e-6
        if active and start is None:
            start = i
        if not active and start is not None:
            rows.append({"period": f"{clock(start)}-{clock(i)}",
                         "emergency_kwh": float(sum(energy[start:i]))})
            start = None
    return rows or [{"period": "无", "emergency_kwh": 0.0}]


def build():
    load, pv, price = read_raw()
    manifest = json.loads((ROOT / "data/results/exp003/evaluation_manifest.json").read_text())
    assert manifest["complete"] is True and manifest["selected_alpha"] == [0.0, 0.0]
    comparisons = []
    for case_id, relative in CASE_FILES.items():
        d = load_dispatch(relative)
        daily = [{"date": day, "dayindex": i + 31, "month": int(day[5:7]),
                  **summary(d, i)} for i, day in enumerate(FORMAL)]
        monthly = []
        for month in range(2, 13):
            indices = [i for i, day in enumerate(FORMAL) if int(day[5:7]) == month]
            monthly.append({"month": month, "days": len(indices), **summary(d, indices)})
        note = "相同的 334 天、固定电价、储能物理边界和结算规则。"
        if case_id == "v2_primary":
            note += "旧模型的共同早停同时使用其他问目标，不能称为严格的第二问输入隔离。"
        if case_id == "exp003_primary":
            note += "一月调参选择两个修正系数均为零；本轮没有证明网络能降低购电费。"
        if case_id == "periodic":
            note += "负载取上周同日，光伏取昨日。当前方案与它的微小费用差别来自数值精度和求解路径，不能解释为模型收益。"
        comparisons.append({"id": case_id, "label": CASE_LABELS[case_id], "seed": 42,
                            "note": note, "days": 334, **summary(d, np.arange(334)),
                            "monthly": monthly, "daily": daily})
    current = next(row for row in comparisons if row["id"] == "exp003_primary")
    expensive = max(current["daily"], key=lambda row: row["total_cost"])["date"]
    selected = list(dict.fromkeys([*SPECIFIED, expensive]))
    d = load_dispatch(CASE_FILES["exp003_primary"])
    details = {}
    for day in selected:
        i = FORMAL.index(day)
        details[day] = {
            "summary": current["daily"][i],
            "intervals": {
                **{dest: d[source][i].tolist() for dest, source in (
                    ("planned_kwh", "original"), ("emergency_kwh", "emergency"),
                    ("charge_kwh", "charge"), ("discharge_kwh", "discharge"),
                    ("states_kwh", "states"))},
                "planned_cost": d["fees"][i, :, 0].tolist(),
                "emergency_cost": d["fees"][i, :, 3].tolist(),
            },
            "table1": [{"period": f"{hour}:00-{hour}:10",
                        "planned_kwh": float(d["original"][i, hour * 6]),
                        "emergency_kwh": float(d["emergency"][i, hour * 6])}
                       for hour in (10, 12, 14, 16, 18, 20)],
            "table2": [{"period": f"{clock(block * 24)}-{clock((block + 1) * 24)}",
                        "charge_kwh": float(d["charge"][i, block * 24:(block + 1) * 24].sum()),
                        "discharge_kwh": float(d["discharge"][i, block * 24:(block + 1) * 24].sum())}
                       for block in range(6)],
            "table3": emergencies(d["emergency"][i]),
        }
    monthly_raw = []
    for month in range(1, 13):
        mask = np.asarray([int(day[5:7]) == month for day in DATES])
        monthly_raw.append({"month": month, "days": int(mask.sum()),
                            "load_kwh": float(load[mask].sum() / 6),
                            "pv_kwh": float(pv[mask].sum() / 6)})
    source_paths = ["C题/C题.md", "data/raw/附件1.csv", "data/raw/附件2_小区负载.csv",
                    "data/raw/附件2_光伏发电实际功率.csv", "data/raw/附件1.xlsx", "data/raw/附件2.xlsx",
                    "data/results/exp003/evaluation_manifest.json", "data/results/exp003/alpha_calibration.json",
                    "experiments/problem2/exp003/protocol.json", "data/results/exp002/protocol.json",
                    "data/results/exp003/dispatch_metrics.csv", "data/results/exp003/daily_metrics.csv",
                    "data/results/exp003/baseline/exp002_q2_costs.csv",
                    "data/results/exp003/baseline/exp002_q2_daily.csv", "data/results/exp003/warmup_2.npz",
                    "data/results/exp002/result2.xlsx", "data/results/exp003/result2.xlsx",
                    *CASE_FILES.values()]
    return {
        "meta": {
            "schema_version": 1, "scope": "第二问人工讨论：既有结果快照，不启动新实验",
            "raw_period": {"start": DATES[0], "end": DATES[-1], "days": 365},
            "formal_period": {"start": FORMAL[0], "end": FORMAL[-1], "days": 334},
            "interval_minutes": 10, "intervals_per_day": 144,
            "units": {"power": "kW", "energy": "kWh", "stored_energy": "kWh",
                      "price": "元/kWh", "cost": "元"},
            "time_convention": "原附件标签是每个区间的终点；00:10 对应 00:00–00:10，0:00+1 对应 23:50–24:00。功率按该十分钟区间平均功率使用，电量=功率÷6。",
            "objective": "总购电费=全部计划购电量×对应固定电价之和+紧急购电量×该时段电价×5之和。计划未用完仍按计划付费。",
            "quantity_convention": "计划购电量与紧急购电量分别展示；如需实际向外网购入的总电量，两者相加。充放电量以储能设备外部交流侧为口径，SOC是设备内储电量。",
            "causality": "每天 00:00 固定当天计划。实际负载和光伏曲线是事后已知的输入，不能作为同日 00:00 已知信息。第二问不用附件3预报和附件4电价。",
            "january": "一月用于历史训练、共同调参和因果暖机；正式费用是附件5要求的 2 月 1 日至 12 月 31 日，不是365天费用。",
            "physics": {"initial_soc_jan1_kwh": 6000, "soc_min_kwh": 1200,
                        "soc_max_kwh": 10800, "power_limit_kw": 5000,
                        "eta_charge": float(np.sqrt(0.9)), "eta_discharge": float(np.sqrt(0.9)),
                        "terminal": "跨日连续；第二问不强加每天首尾相等。"},
            "selected_alpha": [0.0, 0.0], "current_case_id": "exp003_primary",
            "specified_dates": SPECIFIED, "highest_cost_date": expensive,
            "dispatch_detail_dates": selected,
            "dispatch_detail_scope": "仅存四个题目指定日及当前方案总购电费最高日的144区间调度和145个储电量节点；原始输入可查看全年任意日。",
            "comparison_boundary": "四条曲线同物理、同费用口径。v2和exp003仍有训练输入边界差异，费用下降不能全归因于某一项预测技术。当前 α=0 与周期基线基本重合，不能宣称优于周期。",
            "retrospective": "2025年结果已经被团队查看，是现有方案的回顾性比较；本页不会把新的讨论选项写成已验证的改善。",
            "display_metrics": ["总购电费", "计划购电费", "5倍紧急购电费", "计划及紧急购电量", "充放电量", "0时和24时储电量"],
        },
        "tariff": {"interval_starts": [clock(i) for i in range(144)],
                   "interval_ends": [clock(i + 1) for i in range(144)],
                   "price_yuan_per_kwh": price.tolist()},
        "raw": {"dates": DATES, "load_kw": load.tolist(), "pv_kw": pv.tolist(),
                "monthly": monthly_raw,
                "annual": {"load_kwh": float(load.sum() / 6), "pv_kwh": float(pv.sum() / 6)}},
        "comparisons": comparisons, "dispatch_days": details,
        "sources": {path: {"sha256": sha(ROOT / path), "bytes": (ROOT / path).stat().st_size}
                    for path in dict.fromkeys(source_paths)},
    }


def audit_raw_workbooks(evidence):
    workbook = load_workbook(ROOT / "data/raw/附件1.xlsx", read_only=True, data_only=True)
    rows = list(workbook.active.values)
    close([row[1] for row in rows[1:]], evidence["tariff"]["price_yuan_per_kwh"], 0)
    workbook.close()
    workbook = load_workbook(ROOT / "data/raw/附件2.xlsx", read_only=True, data_only=True)
    for sheet, key in (("小区负载", "load_kw"), ("光伏发电实际功率", "pv_kw")):
        rows = list(workbook[sheet].values)[1:]
        assert [row[0].date().isoformat() for row in rows] == evidence["raw"]["dates"]
        close([row[1:] for row in rows], evidence["raw"][key], 0)
    workbook.close()


def audit_result_workbook(relative, d):
    workbook = load_workbook(ROOT / relative, read_only=True, data_only=True)
    sheet = workbook["计划购电量"]
    rows = list(sheet.values)
    assert rows[0][1] == "00:00-00:10" and rows[0][144] == "23:50-24:00"
    assert [row[0].date().isoformat() for row in rows[1:]] == FORMAL
    close([row[1:145] for row in rows[1:]], d["original"])
    close([row[145] for row in rows[1:]], np.sum(d["original"], axis=1))
    close([row[146] for row in rows[1:]], np.sum(d["fees"], axis=(1, 2)))
    rows = list(workbook["充放电量"].values)[1:]
    assert len(rows) == 334 * 6
    for i in range(334):
        for block in range(6):
            row = rows[i * 6 + block]
            close(row[2], d["charge"][i, block * 24:(block + 1) * 24].sum())
            close(row[3], d["discharge"][i, block * 24:(block + 1) * 24].sum())
        close(rows[i * 6][5], d["states"][i, 0])
        close(rows[i * 6 + 1][5], d["states"][i, -1])
    expected = []
    for i, day in enumerate(FORMAL):
        expected.extend((day if j == 0 else None, row["period"], row["emergency_kwh"])
                        for j, row in enumerate(emergencies(d["emergency"][i])))
    actual = list(workbook["紧急购电量"].values)[1:]
    assert len(actual) == len(expected)
    for row, (day, period, energy) in zip(actual, expected):
        assert (row[0].date().isoformat() if row[0] else None) == day
        assert row[1] == period
        close(row[2], energy)
    for sheet in workbook:
        assert not any(cell.data_type == "e" for row in sheet for cell in row)
    workbook.close()


def audit(evidence):
    load, pv, price = read_raw()
    assert evidence["raw"]["dates"] == DATES
    assert evidence["meta"]["raw_period"]["days"] == 365
    assert evidence["meta"]["formal_period"]["days"] == 334
    close(evidence["raw"]["load_kw"], load, 0)
    close(evidence["raw"]["pv_kw"], pv, 0)
    close(evidence["tariff"]["price_yuan_per_kwh"], price, 0)
    audit_raw_workbooks(evidence)
    for key, values in (("load_kwh", load), ("pv_kwh", pv)):
        close(evidence["raw"]["annual"][key], sum(values.ravel()) / 6)
        for row in evidence["raw"]["monthly"]:
            indices = [i for i, day in enumerate(DATES) if int(day[5:7]) == row["month"]]
            assert row["days"] == len(indices)
            close(row[key], sum(values[indices].ravel()) / 6)
        close(sum(row[key] for row in evidence["raw"]["monthly"]), evidence["raw"]["annual"][key])
    metric_rows = csv_rows("data/results/exp003/dispatch_metrics.csv")
    old_metrics = csv_rows("data/results/exp003/baseline/exp002_q2_costs.csv")
    daily_rows = csv_rows("data/results/exp003/daily_metrics.csv")
    old_daily = csv_rows("data/results/exp003/baseline/exp002_q2_daily.csv")
    for case in evidence["comparisons"]:
        d = load_dispatch(CASE_FILES[case["id"]])
        close(d["original"], d["final"], 0)
        close(d["price"], np.tile(price, (334, 1)), 0)
        close(d["actual"][:, :, :2], np.stack((load[31:], pv[31:]), axis=-1), 0)
        close(d["fees"][:, :, 0], d["original"] * price)
        close(d["fees"][:, :, 3], 5 * d["emergency"] * price)
        close(d["fees"][:, :, 1:3], 0, 0)
        close(np.diff(d["states"], axis=1), np.sqrt(.9) * d["charge"] - d["discharge"] / np.sqrt(.9))
        close(d["original"] + pv[31:] / 6 + d["discharge"] + d["emergency"],
              load[31:] / 6 + d["charge"] + d["surplus"])
        close(d["states"][:-1, -1], d["states"][1:, 0])
        assert d["states"].min() >= 1200 - 1e-6 and d["states"].max() <= 10800 + 1e-6
        assert max(d["charge"].max(), d["discharge"].max()) <= 5000 / 6 + 1e-6
        assert not ((d["charge"] > 1e-6) & (d["discharge"] > 1e-6)).any()
        case_name = {"v2_primary": "primary", "exp003_primary": "primary", "periodic": "periodic",
                     "exp003_uncalibrated": "uncalibrated"}[case["id"]]
        source_metrics, source_daily = (old_metrics, old_daily) if case["id"] == "v2_primary" else (metric_rows, daily_rows)
        source = next(row for row in source_metrics if row["name"] == case_name and row["seed"] == "42")
        source_days = [row for row in source_daily if row["name"] == case_name and row["seed"] == "42"]
        assert len(source_days) == 334 and [row["date"] for row in case["daily"]] == FORMAL
        for field in ADDITIVE:
            close(case[field], float(source[field]))
            close(sum(row[field] for row in case["daily"]), case[field])
            close(sum(row[field] for row in case["monthly"]), case[field])
        for i, row in enumerate(case["daily"]):
            assert row["date"] == source_days[i]["date"]
            for field in (*ADDITIVE, "initial_soc", "final_soc"):
                close(row[field], float(source_days[i][field]))
            close(row["planned_cost"] + row["emergency_cost"], row["total_cost"])
        for row in case["monthly"]:
            days = [day for day in case["daily"] if day["month"] == row["month"]]
            assert row["days"] == len(days)
            for field in ADDITIVE:
                close(row[field], sum(day[field] for day in days))
            close(row["initial_soc"], days[0]["initial_soc"])
            close(row["final_soc"], days[-1]["final_soc"])
        if case["id"] in ("v2_primary", "exp003_primary"):
            experiment = "exp002" if case["id"] == "v2_primary" else "exp003"
            audit_result_workbook(f"data/results/{experiment}/result2.xlsx", d)
    warm = load_dispatch("data/results/exp003/warmup_2.npz")
    close(warm["states"][0, 0], 6000, 0)
    for case in evidence["comparisons"]:
        close(case["initial_soc"], warm["states"][-1, -1])
    current = load_dispatch(CASE_FILES["exp003_primary"])
    assert set(evidence["dispatch_days"]) == {*SPECIFIED, evidence["meta"]["highest_cost_date"]}
    for day, detail in evidence["dispatch_days"].items():
        i = FORMAL.index(day)
        fields = {"planned_kwh": "original", "emergency_kwh": "emergency", "charge_kwh": "charge",
                  "discharge_kwh": "discharge", "states_kwh": "states"}
        for key, source in fields.items():
            close(detail["intervals"][key], current[source][i], 0)
        for key, channel in (("planned_cost", 0), ("emergency_cost", 3)):
            close(detail["intervals"][key], current["fees"][i, :, channel], 0)
            close(sum(detail["intervals"][key]), detail["summary"][key])
        for row, hour in zip(detail["table1"], (10, 12, 14, 16, 18, 20)):
            close(row["planned_kwh"], current["original"][i, hour * 6], 0)
            close(row["emergency_kwh"], current["emergency"][i, hour * 6], 0)
        for block, row in enumerate(detail["table2"]):
            for key, source in (("charge_kwh", "charge"), ("discharge_kwh", "discharge")):
                close(row[key], sum(current[source][i, block * 24:(block + 1) * 24]))
        assert detail["table3"] == emergencies(current["emergency"][i])
        close(sum(row["emergency_kwh"] for row in detail["table3"]), detail["summary"]["emergency_kwh"])
    for relative, source in evidence["sources"].items():
        assert sha(ROOT / relative) == source["sha256"]
    return {
        "status": "passed", "raw_days": 365, "formal_days_per_case": 334,
        "compared_cases": len(evidence["comparisons"]), "detailed_dispatch_days": len(evidence["dispatch_days"]),
        "checks": ["附件1、附件2原始XLSX与CSV/JSON全部数值一致", "365天输入和334天结果边界分离",
                   "全年与逐月输入电量=十分钟功率之和÷6", "四方案逐日/逐月/总费用和电量与存档CSV/NPZ一致",
                   "固定电价与5倍紧急结算独立重算一致", "四方案供需平衡/充放电限制/储电量递推和跨日连续满足",
                   "四指定日与最贵日表1/表2/表3和144区间/145储电量一致",
                   "v2与exp003结果XLSX三张工作表全部数值回读一致", "源文件SHA256未变化"],
        "source_count": len(evidence["sources"]), "numeric_tolerance": {"csv_xlsx_raw": 0, "cost_energy": 1e-6},
        "builder_sha256": sha(__file__), "evidence_sha256": sha(HERE / "evidence.json"),
        "authoring_policy": "仅读取已有结果；未重训、未调参、未运行新调度；未修改旧结果。",
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--verify-only", action="store_true")
    args = parser.parse_args()
    output = HERE / "evidence.json"
    if not args.verify_only:
        output.write_text(json.dumps(build(), ensure_ascii=False, allow_nan=False, separators=(",", ":")) + "\n")
    evidence = json.loads(output.read_text())
    result = audit(evidence)
    (HERE / "evidence-audit.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({"status": result["status"], "bytes": output.stat().st_size,
                      "costs": {row["id"]: row["total_cost"] for row in evidence["comparisons"]}}, ensure_ascii=False))


if __name__ == "__main__":
    main()
