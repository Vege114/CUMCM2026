"""Independently verify the exp009 delivery without training or optimization.

Run from any directory: python scripts/verify_exp009_delivery.py
Only delivery_verification.json is written. All calculations below are separate
from the model authors' verification/export implementations. This is a fast
integrity and arithmetic check, not another full-year optimization experiment.
"""
import ast
import hashlib
import json
import sys
from datetime import datetime, timedelta, timezone
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'data/results/exp009'
ETA = np.sqrt(.9)
TOL = 1e-6
SCENARIOS = {'2': 'q2', '3': 'q3', '4-2': 'q4_2', '4-3': 'q4_3'}


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_json(path):
    return json.loads(Path(path).read_text())


def require(condition, message):
    if not condition:
        raise ValueError(message)


def close(actual, expected, name, tolerance=TOL):
    actual, expected = np.asarray(actual, float), np.asarray(expected, float)
    require(actual.shape == expected.shape, f'{name}: shape {actual.shape} != {expected.shape}')
    require(np.isfinite(actual).all() and np.isfinite(expected).all(), f'{name}: nonfinite')
    error = float(np.max(np.abs(actual-expected), initial=0.))
    require(error <= tolerance, f'{name}: maximum error {error} > {tolerance}')
    return error


def load_arrays(path):
    with np.load(path, allow_pickle=False) as archive:
        return {key: archive[key].copy() for key in archive.files}


def inputs():
    raw = ROOT/'data/raw'
    actual = []
    for name in ('附件2_小区负载.csv', '附件2_光伏发电实际功率.csv', '附件4.csv'):
        frame = pd.read_csv(raw/name)
        require(frame.shape == (365, 145), f'{name}: dimensions')
        require(pd.DatetimeIndex(pd.to_datetime(frame.iloc[:, 0])).equals(
            pd.date_range('2025-01-01', '2025-12-31')), f'{name}: dates')
        actual.append(frame.iloc[:, 1:].to_numpy(float))
    reference = pd.read_csv(raw/'附件1.csv').iloc[:, 1:].to_numpy(float)
    require(reference.shape == (144, 3), 'attachment1 dimensions')
    return np.stack(actual, axis=-1), reference


def verify_manifest():
    manifest = read_json(OUT/'复现清单.json')
    require(manifest['run_days'] == 365 and manifest['evaluation_days'] == 334, 'manifest date scope')
    require(manifest['evaluation_dates'] == ['2025-02-01', '2025-12-31'], 'manifest evaluation dates')
    bindings = manifest['source_bindings']
    require(len(bindings) > 10, 'source dependency bindings missing')
    sources = set()
    for item in bindings:
        source, copy = ROOT/item['source'], ROOT/item['paper_copy']
        require(digest(source) == item['sha256'], f'source hash: {source}')
        require(digest(copy) == item['sha256'], f'paper-copy hash: {copy}')
        require(source.stat().st_size == item['bytes'], f'source size: {source}')
        sources.add(item['source'])
    require({'experiments/exp008/q1.py', 'experiments/exp009/q12_run.py',
             'experiments/exp009/q34_run.py'}.issubset(sources), 'model entry bindings')
    counts = {'source_bindings': len(bindings)}
    for category in ('inputs', 'trajectories'):
        require(bool(manifest[category]), f'empty manifest {category}')
        for path, expected in manifest[category].items():
            require(digest(ROOT/path) == expected, f'{category} hash: {path}')
        counts[category] = len(manifest[category])
    required_q1 = {f'data/results/exp009/q1/{name}.json' for name in ('baseline', 'revised')}
    require(required_q1.issubset(manifest['trajectories']), 'manifest must bind Q1 trajectories')
    required_workbooks = {f'result{name}.xlsx' for name in ('1', '2', '3', '4-2', '4-3')}
    require(set(manifest['workbooks']) == required_workbooks, 'manifest five workbooks')
    for name, expected in manifest['workbooks'].items():
        require(digest(OUT/name) == expected, f'workbook hash: {name}')
    counts['workbooks'] = 5
    return counts


def verify_protocol(directory, scenario):
    protocol = read_json(directory/'protocol.json')
    hashes = protocol.get('source_sha256', protocol.get('sources', {}))
    require(bool(hashes), 'run source hashes missing')
    documented_changes = []
    for path, expected in hashes.items():
        current = digest(ROOT/path)
        if current != expected:
            record_path = OUT/'source_documentation_changes.json'
            require(record_path.is_file(), f'run protocol source changed: {path}')
            matches = [item for item in read_json(record_path)['changes'] if item['source'] == path]
            require(len(matches) == 1, f'unrecorded/ambiguous source change: {path}')
            item = matches[0]
            before = ROOT/item['before_source']
            require(digest(before) == expected == item['before_sha256'], f'run source snapshot hash: {path}')
            require(current == item['after_sha256'], f'documented current source hash: {path}')
            trees = []
            for source in (before, ROOT/path):
                tree = ast.parse(source.read_text())
                if (tree.body and isinstance(tree.body[0], ast.Expr)
                        and isinstance(tree.body[0].value, ast.Constant)
                        and isinstance(tree.body[0].value.value, str)):
                    tree.body.pop(0)
                trees.append(ast.dump(tree, include_attributes=False))
            require(trees[0] == trees[1], f'executable AST changed: {path}')
            documented_changes.append({'source': path, 'run_sha256': expected,
                'current_sha256': current, 'executable_ast_identical': True, 'reason': item['reason']})
    raw = protocol.get('raw_input_sha256', protocol.get('raw_data_sha256', {}))
    require(len(raw) >= 5, 'run raw hashes missing')
    for name, expected in raw.items():
        require(digest(ROOT/'data/raw'/name) == expected, f'run input changed: {name}')
    forecast_path = protocol.get('forecast_path',
        'data/results/exp008/forecast_absolute_hgb/direct_hgb_ridge28_memory_half.npz')
    forecast_hash = protocol.get('forecast_sha256', protocol.get('forecast_archive_sha256'))
    require(digest(ROOT/forecast_path) == forecast_hash, 'frozen issued forecast hash')
    require(protocol['scenario'] == scenario, 'protocol scenario')
    return {'source_hashes': len(hashes), 'raw_hashes': len(raw), 'forecast_hash': forecast_hash,
            'documented_nonexecutable_changes': documented_changes}


def verify_annual(scenario, actual, reference):
    directory = OUT/SCENARIOS[scenario]
    path = directory/('dispatch.npz' if scenario in ('2', '4-2') else 'dispatch_365.npz')
    a = load_arrays(path)
    require(np.array_equal(a['days'], np.arange(365)), f'{scenario}: complete 365 day IDs')
    for key in ('original', 'final', 'price', 'charge', 'discharge', 'emergency', 'surplus'):
        require(a[key].shape == (365, 144), f'{scenario}/{key}: shape')
        require(np.isfinite(a[key]).all(), f'{scenario}/{key}: finite')
        require(a[key].min() >= -TOL, f'{scenario}/{key}: nonnegative')
    require(a['states'].shape == (365, 145), f'{scenario}: state dimensions')
    close(a['actual'], actual, f'{scenario}: actual inputs', 0.)
    expected_price = np.broadcast_to(reference[:, 0], (365, 144)) if scenario in ('2', '3') else actual[:, :, 2]
    close(a['price'], expected_price, f'{scenario}: settlement prices', 0.)
    g, q, c, d, e, w, s, p = (a[key] for key in
        ('original', 'final', 'charge', 'discharge', 'emergency', 'surplus', 'states', 'price'))
    errors = {
        'balance_kwh': close(q+(actual[:, :, 1]-actual[:, :, 0])/6+d+e-c-w,
                             np.zeros((365, 144)), f'{scenario}: energy balance'),
        'state_kwh': close(np.diff(s, axis=1), ETA*c-d/ETA, f'{scenario}: SOC recurrence'),
        'cross_midnight_kwh': close(s[1:, 0], s[:-1, -1], f'{scenario}: midnight continuity')}
    close(s[0, 0], 6000., f'{scenario}: Jan1 SOC')
    require(s.min() >= 1200-TOL and s.max() <= 10800+TOL, f'{scenario}: SOC bounds')
    require(max(c.max(), d.max()) <= 5000/6+TOL, f'{scenario}: power bounds')
    require(not ((c>TOL)&(d>TOL)).any(), f'{scenario}: simultaneous charge/discharge')
    require(not ((c>TOL)&(e>TOL)).any(), f'{scenario}: emergency charging')
    bill = np.stack((g*p, 1.5*p*np.maximum(q-g, 0), -.5*p*np.maximum(g-q, 0), 5*p*e), axis=-1)
    errors['slot_fee_yuan'] = close(a['fees'], bill, f'{scenario}: actual signed refund billing')
    summary = read_json(directory/'summary.json')
    if scenario in ('2', '4-2'):
        close(q, g, f'{scenario}: fixed midnight purchase')
        mask = a['allowed_charge']
        require(mask.shape == (365, 144), f'{scenario}: mask dimensions')
        close(mask, np.round(mask), f'{scenario}: binary mode mask')
        require(mask.min() >= 0 and mask.max() <= 1, f'{scenario}: mode mask bounds')
        require(not ((c>TOL)&(mask<.5)).any(), f'{scenario}: charging outside allowed mode')
        require(not ((d>TOL)&(mask>.5)).any(), f'{scenario}: discharging outside allowed mode')
        if scenario == '2':
            prior = np.r_[1., mask[:-1, -1]]
            changes = (np.diff(mask, axis=1) != 0).sum(axis=1)+(mask[:, 0] != prior)
            require(changes.max() <= 8, 'Q2 daily mode changes including midnight >8')
        else:
            close(mask.reshape(365, 24, 6), np.repeat(mask[:, ::6, None], 6, axis=-1),
                  'Q4-2 hourly shared modes')
        require(summary['complete'] and summary['days'] == 365, f'{scenario}: complete flag')
        close(summary['full_year_cost'], bill.sum(), f'{scenario}: full-year summary', 1e-5)
        close(summary['feb_dec_cost'], bill[31:].sum(), f'{scenario}: 334-day summary', 1e-5)
        close(summary['feb_initial_soc'], s[31, 0], f'{scenario}: Feb1 state')
        audits = read_json(directory/'planning_audit.json')
        require(len(audits) == 365, f'{scenario}: 365 planning records')
        for day, audit in enumerate(audits):
            require(audit['day'] == day and audit['information_cutoff'] == day*144,
                    f'{scenario}: issue cutoff day {day}')
            require(audit['purchase_locked_before_current_actual_read'], f'{scenario}: locked purchase')
            require(audit['mip']['feasible'], f'{scenario}: MIP feasibility day {day}')
            if day:
                require(audit['boundary_before'] == audits[day-1]['boundary_after'],
                        f'{scenario}: planned mode boundary day {day}')
            for section in ('forecast', 'history'):
                index = audit[section].get('max_observed_index')
                require(index is None or index < day*144, f'{scenario}: future data in {section}')
    else:
        require(a['versions'].shape == (365, 4, 144), f'{scenario}: legal issue versions')
        close(g, a['versions'][:, 0], f'{scenario}: midnight version')
        chosen = np.concatenate([a['versions'][:, j, j*36:(j+1)*36] for j in range(4)], axis=1)
        close(q, chosen, f'{scenario}: last legally released version')
        for j in range(1, 4):
            close(a['versions'][:, j, :j*36], q[:, :j*36], f'{scenario}: immutable executed prefix')
        formal = load_arrays(directory/f'dispatch_{scenario}.npz')
        warmup = load_arrays(directory/'warmup.npz')
        for key in a:
            close(formal[key], a[key][31:], f'{scenario}: 334-day slice {key}', 0.)
            close(warmup[key], a[key][:31], f'{scenario}: January slice {key}', 0.)
        require(summary['days'] == 334, f'{scenario}: summary days')
        close(summary['total_cost'], bill[31:].sum(), f'{scenario}: 334-day summary', 1e-5)
        close(summary['initial_soc_february1_kwh'], s[31, 0], f'{scenario}: Feb1 state')
        audits = read_json(directory/'audit.json')
        require(len(audits) == 1460, f'{scenario}: 1460 release records')
        for i, audit in enumerate(audits):
            day, slot = divmod(i, 4)[0], (i % 4)*36
            require((audit['day'], audit['slot']) == (day, slot), f'{scenario}: audit issue order')
            require(audit['solver']['feasible'], f'{scenario}: LP feasibility')
            require(all(stop <= day*144 for stop in audit['risk']['label_stops_exclusive']),
                    f'{scenario}: incomplete historical target day')
            index = audit['forecast'].get('max_observed_index')
            require(index is None or index < day*144+slot, f'{scenario}: future observation')
    close(summary['january_cost'], bill[:31].sum(), f'{scenario}: January summary', 1e-5)
    return a, {'days': 365, 'slots': 52560, 'evaluation_days': 334, 'errors': errors,
               'full_year_cost_yuan': float(bill.sum()), 'feb_dec_cost_yuan': float(bill[31:].sum()),
               'feb_initial_soc_kwh': float(s[31, 0]), 'planning_records': len(audits),
               'protocol': verify_protocol(directory, scenario)}


def clock(slot):
    return f'{slot//6:02d}:{slot%6*10:02d}'


def interval(start, stop):
    return f'{clock(start)}-{clock(stop)}'


def normalize_time(value):
    # Preserve meaning while accepting the template's unpadded clock text.
    parts = str(value).split('-')
    normalized = []
    for part in parts:
        h, m = part.split(':')[:2]
        normalized.append(f'{int(h):02d}:{int(m):02d}')
    return '-'.join(normalized)


def verify_workbook(scenario, a):
    path = OUT/f'result{scenario}.xlsx'
    workbook = load_workbook(path, data_only=False)
    names = ['计划购电量']+(['调整购电量'] if scenario in ('3', '4-3') else [])+['充放电量', '紧急购电量']
    require(workbook.sheetnames == names, f'{path.name}: sheet names')
    require(not any(cell.data_type == 'f' for sheet in workbook for row in sheet for cell in row),
            f'{path.name}: unexpected unevaluated formulas')
    a = {key: value[31:] for key, value in a.items()}
    dates = [datetime(2025, 2, 1)+timedelta(days=i) for i in range(334)]
    cells, maximum = 0, 0.
    for name, key in [('计划购电量', 'original')]+([('调整购电量', 'final')] if scenario in ('3', '4-3') else []):
        sheet = workbook[name]
        require((sheet.max_row, sheet.max_column) == (335, 147), f'{path.name}/{name}: dimensions')
        for slot in range(144):
            require(normalize_time(sheet.cell(1, slot+2).value) == interval(slot, slot+1),
                    f'{path.name}/{name}: interval column {slot}')
        for i, date in enumerate(dates):
            row = list(sheet.iter_rows(min_row=i+2, max_row=i+2, values_only=True))[0]
            require(row[0] == date, f'{path.name}/{name}: date {i}')
            fee = a['fees'][i, :, 0].sum() if name == '计划购电量' and scenario in ('3', '4-3') else a['fees'][i].sum()
            wanted = np.r_[a[key][i], a[key][i].sum(), fee]
            maximum = max(maximum, close(row[1:], wanted, f'{path.name}/{name}: row {i}'))
            cells += len(wanted)
    sheet = workbook['充放电量']
    require(sheet.max_row == 2005 and sheet.max_column == 6, f'{path.name}: storage dimensions')
    for i, date in enumerate(dates):
        for block in range(6):
            row = list(sheet.iter_rows(min_row=2+6*i+block, max_row=2+6*i+block, values_only=True))[0]
            require(row[0] == (date if block == 0 else None), f'{path.name}: storage date')
            require(normalize_time(row[1]) == interval(block*24, (block+1)*24), f'{path.name}: storage interval')
            expected = [a[key][i, block*24:(block+1)*24].sum() for key in ('charge', 'discharge')]
            maximum = max(maximum, close(row[2:4], expected, f'{path.name}: storage quantities'))
            cells += 2
            if block < 2:
                require(normalize_time(row[4]) == ('00:00' if block == 0 else '24:00'), f'{path.name}: SOC label')
                maximum = max(maximum, close(row[5], a['states'][i, 0 if block == 0 else -1], f'{path.name}: endpoint SOC'))
                cells += 1
            else:
                require(row[4] is None and row[5] is None, f'{path.name}: unexpected SOC fields')
    expected_rows = []
    for i, date in enumerate(dates):
        active = a['emergency'][i] > TOL
        edges = np.diff(np.r_[False, active, False].astype(int))
        starts, stops = np.flatnonzero(edges == 1), np.flatnonzero(edges == -1)
        if not len(starts):
            expected_rows.append((date, '无', 0.))
        for j, (start, stop) in enumerate(zip(starts, stops)):
            expected_rows.append((date if j == 0 else None, interval(start, stop), a['emergency'][i, start:stop].sum()))
    sheet = workbook['紧急购电量']
    got_rows = [row for row in sheet.iter_rows(min_row=2, values_only=True) if any(v is not None for v in row)]
    require(len(got_rows) == len(expected_rows), f'{path.name}: emergency row count')
    for got, expected in zip(got_rows, expected_rows):
        require(got[0] == expected[0], f'{path.name}: emergency date')
        label = got[1] if got[1] == '无' else normalize_time(got[1])
        require(label == expected[1], f'{path.name}: emergency contiguous interval')
        maximum = max(maximum, close(got[2], expected[2], f'{path.name}: emergency quantity'))
        cells += 1
    workbook.close()
    return {'numeric_cells_checked': cells, 'max_numeric_error': maximum, 'days': 334,
            'emergency_rows': len(expected_rows), 'sha256': digest(path), 'formulas': 0}


def verify_q1(reference):
    case = read_json(OUT/'q1/revised.json')
    a = {key: np.asarray(value, float) for key, value in case['trajectory'].items()}
    g, c, d, w, s = (a[key] for key in ('g', 'c', 'd', 'w', 'E'))
    require(g.shape == (144,) and s.shape == (145,), 'Q1 dimensions')
    close(g+(reference[:, 2]-reference[:, 1]+d-c)/6-w, np.zeros(144), 'Q1 balance')
    close(np.diff(s), (ETA*c-d/ETA)/6, 'Q1 SOC recurrence')
    close(s[[0, -1]], [6000., 6000.], 'Q1 endpoint SOC')
    require(s.min() >= 1200-TOL and s.max() <= 10800+TOL, 'Q1 SOC bounds')
    require(min(g.min(), c.min(), d.min(), w.min()) >= -TOL, 'Q1 nonnegative')
    require(max(c.max(), d.max()) <= 5000+TOL, 'Q1 power bound')
    require(not ((c>TOL)&(d>TOL)).any(), 'Q1 simultaneous modes')
    power = c-d
    require(np.max(np.abs(power-np.roll(power, 1))) <= 1000+TOL, 'Q1 circular ramp')
    for direction, values in [('c', c), ('d', d)]:
        mode = np.rint(a['z'+direction]).astype(int)
        require(np.all(values >= mode-TOL) and np.all(values <= 5000*mode+TOL), 'Q1 min on power')
        for i in range(144):
            if mode[i] and not mode[(i-1) % 144]:
                require(all(mode[(i+j) % 144] for j in range(3)), 'Q1 min on 3')
            if not mode[i] and mode[(i-1) % 144]:
                require(all(not mode[(i+j) % 144] for j in range(2)), 'Q1 min off 2')
    cost = float(g @ reference[:, 0])
    close(case['metrics']['cost'], cost, 'Q1 cost')
    require(len(case['stages']) == 2, 'Q1 must be accepted two-layer model')
    path = OUT/'result1.xlsx'
    workbook = load_workbook(path, data_only=False)
    require(workbook.sheetnames == ['计划购电量', '充放电量'], 'Q1 workbook sheets')
    require(not any(cell.data_type == 'f' for sheet in workbook for row in sheet for cell in row), 'Q1 formulas')
    sheet = workbook['计划购电量']
    require((sheet.max_row, sheet.max_column) == (145, 2), 'Q1 purchase dimensions')
    for i in range(144):
        require(normalize_time(sheet.cell(i+2, 1).value) == interval(i, i+1), 'Q1 interval alignment')
        close(sheet.cell(i+2, 2).value, g[i], 'Q1 workbook purchase')
    sheet = workbook['充放电量']
    require((sheet.max_row, sheet.max_column) == (7, 5), 'Q1 battery dimensions')
    for block in range(6):
        row = list(sheet.iter_rows(min_row=block+2, max_row=block+2, values_only=True))[0]
        require(normalize_time(row[0]) == interval(block*24, (block+1)*24), 'Q1 storage interval')
        close(row[1:3], [x[block*24:(block+1)*24].sum()/6 for x in (c, d)], 'Q1 workbook storage')
        if block < 2:
            require(normalize_time(row[3]) == ('00:00' if block == 0 else '24:00'), 'Q1 SOC label')
            close(row[4], s[0 if block == 0 else -1], 'Q1 workbook SOC')
        else:
            require(row[3] is None and row[4] is None, 'Q1 unexpected endpoint fields')
    workbook.close()
    return {'cost_yuan': cost, 'slots': 144, 'numeric_workbook_cells': 158,
            'workbook_sha256': digest(path), 'accepted_layers': 2}


def main():
    report = {'passed': False, 'scope': 'Independent artifact integrity, all-slot physics and actual billing, all five value workbooks; no training or optimization.',
              'generated_at': datetime.now(timezone.utc).isoformat(), 'checks': {}, 'errors': []}
    actual, reference = inputs()
    tasks = [('manifest', verify_manifest), ('q1', lambda: verify_q1(reference))]
    for name, task in tasks:
        try:
            report['checks'][name] = task()
        except Exception as error:
            report['errors'].append(f'{name}: {type(error).__name__}: {error}')
    for scenario in SCENARIOS:
        try:
            arrays, result = verify_annual(scenario, actual, reference)
            report['checks'][scenario] = result
            report['checks'][f'workbook_{scenario}'] = verify_workbook(scenario, arrays)
        except Exception as error:
            report['errors'].append(f'{scenario}: {type(error).__name__}: {error}')
    report['passed'] = not report['errors']
    destination = OUT/'delivery_verification.json'
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False)+'\n')
    print(json.dumps({'passed': report['passed'], 'errors': report['errors'],
                      'report': str(destination.relative_to(ROOT))}, ensure_ascii=False))
    return 0 if report['passed'] else 1


if __name__ == '__main__':
    sys.exit(main())
