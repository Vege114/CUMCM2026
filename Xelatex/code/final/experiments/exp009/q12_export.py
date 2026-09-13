"""Export Q1, Q2 and Q4-2 into the supplied result workbook templates.

The year-long controls start January 1; annual result workbooks contain the
February--December range required by attachment 5. Values are calculated before
writing, and the original value-only templates require no formula recalculation.
"""
import argparse
from copy import copy
import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'data/results/exp009'


def clock(slot):
    return f'{slot//6:02d}:{slot%6*10:02d}'


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def export(scenario):
    directory=OUT/({'1':'q1','2':'q2','4-2':'q4_2'}[scenario])
    template=ROOT/f'data/templates/result{scenario}.xlsx'
    book=load_workbook(template)
    names=['计划购电量','充放电量']+([] if scenario=='1' else ['紧急购电量'])
    assert book.sheetnames==names
    assert not any(c.data_type=='f' for s in book for row in s for c in row)
    expected=[]
    def put(sheet,row,column,value,style=None):
        cell=book[sheet].cell(row,column,value)
        if style is not None:cell._style=copy(style)
        if isinstance(value,(float,np.floating)):
            cell.number_format='0.00' if column==147 else '0.0000'
        elif isinstance(value,pd.Timestamp):
            cell.number_format='yyyy/mm/dd'
        expected.append((sheet,row,column,value))
    if scenario=='1':
        archive=directory/'revised.json'
        result=json.loads(archive.read_text())
        assert len(result['stages'])==2
        q={k:np.asarray(v) for k,v in result['trajectory'].items()}
        for t in range(144):
            put('计划购电量',t+2,1,f'{clock(t)}-{clock(t+1)}')
            put('计划购电量',t+2,2,float(q['g'][t]))
        for block in range(6):
            start,stop=24*block,24*(block+1)
            put('充放电量',block+2,1,f'{clock(start)}-{clock(stop)}')
            put('充放电量',block+2,2,max(0.,float(q['c'][start:stop].sum()/6)))
            put('充放电量',block+2,3,max(0.,float(q['d'][start:stop].sum()/6)))
            put('充放电量',block+2,4,'00:00' if block==0 else '24:00' if block==1 else None)
            put('充放电量',block+2,5,float(q['E'][0]) if block==0 else float(q['E'][-1]) if block==1 else None)
        total=result['metrics']['cost']; days=1; emergency_count=0
    else:
        status=json.loads((directory/'summary.json').read_text())
        assert status['complete'] and status['days']==365
        archive=directory/'dispatch.npz'
        with np.load(archive,allow_pickle=False) as z:
            assert np.array_equal(z['days'],np.arange(365))
            a={k:z[k][31:].copy() for k in z.files}
        dates=pd.date_range('2025-02-01','2025-12-31')
        sheet='计划购电量'
        put(sheet,1,1,'日期\\时间')
        for t in range(144):put(sheet,1,t+2,f'{clock(t)}-{clock(t+1)}')
        put(sheet,1,146,'全天计划购电量/kWh')
        put(sheet,1,147,'全天总购电费/元')
        for i,date in enumerate(dates):
            put(sheet,i+2,1,date)
            for t in range(144):put(sheet,i+2,t+2,float(a['original'][i,t]))
            put(sheet,i+2,146,float(a['original'][i].sum()))
            put(sheet,i+2,147,float(a['fees'][i].sum()))
        book[sheet].freeze_panes='B2'
        storage=book['充放电量']; styles=[copy(storage.cell(2,j)._style) for j in range(1,7)]
        for row in storage:
            if row[0].row>1:
                for cell in row:cell.value=None
        for i,date in enumerate(dates):
            for block in range(6):
                start,stop=block*24,(block+1)*24
                row=[date if block==0 else None,f'{clock(start)}-{clock(stop)}',
                    float(a['charge'][i,start:stop].sum()),float(a['discharge'][i,start:stop].sum()),
                    '00:00' if block==0 else '24:00' if block==1 else None,
                    float(a['states'][i,0]) if block==0 else float(a['states'][i,-1]) if block==1 else None]
                for j,value in enumerate(row,1):put('充放电量',2+i*6+block,j,value,styles[j-1])
        book['充放电量'].freeze_panes='C2'
        emergency=book['紧急购电量'];styles=[copy(emergency.cell(2,j)._style) for j in range(1,4)]
        for row in emergency:
            if row[0].row>1:
                for cell in row:cell.value=None
        rows=[]
        for i,date in enumerate(dates):
            edges=np.diff(np.r_[False,a['emergency'][i]>1e-6,False].astype(int))
            starts,stops=np.flatnonzero(edges==1),np.flatnonzero(edges==-1)
            if not len(starts):rows.append([date,'无',0.])
            for j,(start,stop) in enumerate(zip(starts,stops)):
                rows.append([date if j==0 else None,f'{clock(int(start))}-{clock(int(stop))}',
                    float(a['emergency'][i,start:stop].sum())])
        for i,row in enumerate(rows,2):
            for j,value in enumerate(row,1):put('紧急购电量',i,j,value,styles[j-1])
        total=float(a['fees'].sum()); days=334;emergency_count=len(rows)
    output=OUT/f'result{scenario}.xlsx'
    book.save(output)
    restored=load_workbook(output,data_only=False)
    assert restored.sheetnames==names
    max_error=0.; numeric=0
    for sheet,row,col,wanted in expected:
        got=restored[sheet].cell(row,col).value
        if isinstance(wanted,(float,np.floating)):
            assert isinstance(got,(int,float)) and np.isfinite(got)
            max_error=max(max_error,abs(got-wanted)); numeric+=1
        elif isinstance(wanted,pd.Timestamp):assert got==wanted.to_pydatetime()
        else:assert got==wanted,(sheet,row,col,got,wanted)
    assert max_error<1e-6
    if scenario=='1':assert restored['计划购电量'].max_row==145
    else:assert restored['计划购电量'].max_row==335 and restored['计划购电量'].max_column==147
    assert not any(c.data_type=='f' for s in restored for row in s for c in row)
    check={'passed':True,'file':str(output.relative_to(ROOT)),'file_sha256':digest(output),
        'archive_sha256':digest(archive),'template_sha256':digest(template),'days':days,
        'all_written_cells_checked':len(expected),'numeric_cells_checked':numeric,
        'max_numeric_error':max_error,'emergency_rows_checked':emergency_count,'formulas':0,
        'emergency_interval_threshold_kwh':1e-6,'total_cost':total,
        'time_labels':'corrected template one-slot shift to half-open midnight-to-midnight intervals',
        'fee_field':'actual total bill including emergency; Q1 has no fee column in template'}
    (directory/'workbook_verification.json').write_text(json.dumps(check,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(check,ensure_ascii=False))


if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--scenario',choices=['1','2','4-2','all'],default='all')
    args=p.parse_args()
    for scenario in ('1','2','4-2') if args.scenario=='all' else (args.scenario,):export(scenario)
